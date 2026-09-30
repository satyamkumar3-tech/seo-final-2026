#!/usr/bin/env python3
"""Update status files, product rotation, and Excel workbook for Week 8 Rank 158."""
import os
import csv
import json
import openpyxl
from pathlib import Path
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]

RANK = "158"
SLUG = "party-wear-earrings-2026"
PRIMARY = "party wear earrings"
POST_ID = "38850"
URL = "https://blog.bluestone.com/party-wear-earrings-2026/"
TODAY = "2026-09-15"

EXEC_NOTE = (
    "Published 2026-09-15, WP post ID 38850, year 2026, author Satyam (270271337), "
    "Yoast focus KW party wear earrings, schema OK, carousel SKUs: BIJP0686H03 (38844), "
    "BINK0363H03 (38845), BISP0427H21 (38846), BIPM0001H28 (38847), BIPN0880H218 (38848), "
    "BIIP0427H16 (38849); Type 3 SKUs: BIIP0279S08 The Aleena Huggie Earrings (hero 38851), "
    "BIIP0427H16 The Vicky Hoop Earrings (flatlay 38852), BISA0255D05 The Asya Huggie Earrings (lifestyle 38853); "
    "flatlay setting: marble-vanity; passed live QA with 3002 visible words, 0 dashes, 6 verified HTTP 200 carousel cards, and matching og:image."
)

def update_queue_status_csv():
    status_csv = ROOT / "output/Week8_Blog_Queue_status.csv"
    rows = []
    header = ['rank', 'primary', 'slug', 'blog_url', 'wp_post_id', 'status', 'carousel_media', 'type3_media', 'lines', 'visible_words', 'notes']
    if status_csv.exists():
        with open(status_csv, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)
        if rows:
            header = rows[0]
            rows = rows[1:]
            
    updated = False
    new_row = [
        RANK,
        PRIMARY,
        SLUG,
        URL,
        POST_ID,
        "published",
        "38844-38849",
        "hero 38851; flatlay 38852; lifestyle 38853",
        "425",
        "3002",
        EXEC_NOTE
    ]
    
    for i, r in enumerate(rows):
        if r and r[0] == RANK:
            rows[i] = new_row
            updated = True
            break
            
    if not updated:
        rows.append(new_row)
        
    with open(status_csv, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)
    print(f"Updated {status_csv} with Rank {RANK}")

def update_product_rotation():
    rot_file = ROOT / "output/product_rotation.json"
    with open(rot_file, "r", encoding="utf-8") as f:
        rot = json.load(f)
        
    rot["updated"] = TODAY
    
    # Upsert recent flatlay settings
    flatlay_entry = {
        "rank": f"Week8 Rank {RANK}",
        "setting": "marble-vanity",
        "note": "party wear earrings The Vicky Hoop Earrings on marble vanity"
    }
    # Check if already present
    cleaned_flatlay = []
    for e in rot.get("recent_flatlay_settings", []):
        if isinstance(e, dict):
            if e.get("rank") != flatlay_entry["rank"]:
                cleaned_flatlay.append(e)
        else:
            cleaned_flatlay.append(e)
    cleaned_flatlay.append(flatlay_entry)
    rot["recent_flatlay_settings"] = cleaned_flatlay
    
    # Upsert recent type3 trios
    trio_entry = {
        "rank": f"Week8 Rank {RANK}",
        "hero": "BIIP0279S08 The Aleena Huggie Earrings",
        "flatlay": "BIIP0427H16 The Vicky Hoop Earrings",
        "lifestyle": "BISA0255D05 The Asya Huggie Earrings"
    }
    cleaned_trios = []
    for t in rot.get("recent_type3_trios", []):
        if isinstance(t, dict):
            if t.get("rank") != trio_entry["rank"]:
                cleaned_trios.append(t)
        else:
            cleaned_trios.append(t)
    cleaned_trios.append(trio_entry)
    rot["recent_type3_trios"] = cleaned_trios
    
    with open(rot_file, "w", encoding="utf-8") as f:
        json.dump(rot, f, indent=2)
    print(f"Updated {rot_file} with Rank {RANK} flatlay setting and Type 3 trio")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    wb = openpyxl.load_workbook(xlsx_path)
    if "Week 8" not in wb.sheetnames:
        print("Week 8 sheet not found in Excel!")
        return
        
    ws = wb["Week 8"]
    headers = [cell for cell in next(ws.iter_rows(values_only=True))]
    
    col_rank = headers.index("Priority Rank") + 1
    col_action = headers.index("Action") + 1
    col_url = headers.index("Bluestone Blog URL") + 1
    col_note = headers.index("Execution Note") + 1
    
    updated = False
    for row in range(2, ws.max_row + 1):
        cell_val = str(ws.cell(row=row, column=col_rank).value)
        if cell_val == RANK:
            ws.cell(row=row, column=col_action, value="Done")
            ws.cell(row=row, column=col_url, value=URL)
            ws.cell(row=row, column=col_note, value=EXEC_NOTE)
            updated = True
            print(f"Updated row {row} in Week 8 sheet for Rank {RANK}")
            break
            
    if updated:
        wb.save(xlsx_path)
        print(f"Saved {xlsx_path}")
    else:
        print(f"Rank {RANK} row not found in Week 8 sheet!")
    wb.close()

def main():
    update_queue_status_csv()
    update_product_rotation()
    update_excel_workbook()

if __name__ == "__main__":
    main()
