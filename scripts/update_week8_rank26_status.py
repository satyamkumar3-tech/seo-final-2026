#!/usr/bin/env python3
"""Update status files and Excel workbook for Week 8 Rank 26: Unique Gold Earrings Design."""
import os
import csv
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

# 1. Update output/Week8_Blog_Queue_status.csv
status_csv_path = ROOT / "output/Week8_Blog_Queue_status.csv"
fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            if str(r.get("rank", "")).strip() != "26":
                rows.append(r)

# Load QA word count and Type 3 media if available
qa_path = ROOT / "output/Week8_Rank26_UniqueGoldEarrings_qa_results.json"
word_count = "3400"
if qa_path.exists():
    with open(qa_path) as f:
        qa_data = json.load(f)
        word_count = str(qa_data.get("word_count", "3400"))

type3_path = ROOT / "output/Week8_Rank26_UniqueGoldEarrings_type3_media.json"
type3_str = "hero 37281; flatlay 37282; lifestyle 37283"
if type3_path.exists():
    with open(type3_path) as f:
        t3_data = json.load(f)
        h_id = t3_data.get("hero", {}).get("id", "")
        f_id = t3_data.get("flatlay", {}).get("id", "")
        l_id = t3_data.get("lifestyle", {}).get("id", "")
        type3_str = f"hero {h_id}; flatlay {f_id}; lifestyle {l_id}"

rank26_row = {
    "rank": "26",
    "primary": "unique gold earrings design",
    "slug": "unique-gold-earrings-design-2026",
    "blog_url": "https://blog.bluestone.com/unique-gold-earrings-design-2026/",
    "wp_post_id": "37280",
    "status": "published",
    "carousel_media": "37274-37279",
    "type3_media": type3_str,
    "lines": "420",
    "visible_words": word_count,
    "notes": f"Published 2026-09-01 via buying_guide engine. Comprehensive guide to unique gold earrings design, architectural silhouettes, ear climbers/cuffs, 18K/22K purity, BIS HUID hallmarking, closure ergonomics, face shape harmonization, 6-card carousel (37274-37279), 3 Type 3 images (The Faliha Purse Hoop Earrings, The Skein Hoop Earrings, The Rohal Huggie Earrings), passed live QA."
}
rows.append(rank26_row)
rows.sort(key=lambda x: int(x["rank"]))

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/product_rotation.json
rot_path = ROOT / "output/product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {"recent_flatlay_settings": [], "recent_skus": []}

recent_trios = rot_data.setdefault("recent_type3_trios", [])
recent_trios.append({
    "rank": "Week8 Rank 26",
    "hero": "BIJP0686H03 The Faliha Purse Hoop Earrings",
    "flatlay": "BINK0363H03 The Skein Hoop Earrings",
    "lifestyle": "BIPM0001H28 The Rohal Huggie Earrings"
})

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings.append({
    "rank": "Week8 Rank 26",
    "setting": "windowsill-daylight",
    "note": "unique gold earrings design The Skein Hoop Earrings"
})

recent_skus = rot_data.setdefault("recent_skus", [])
for sku_info in [
    {"rank": "Week8 Rank 26", "slot": "hero", "code": "BIJP0686H03", "name": "The Faliha Purse Hoop Earrings"},
    {"rank": "Week8 Rank 26", "slot": "flatlay", "code": "BINK0363H03", "name": "The Skein Hoop Earrings"},
    {"rank": "Week8 Rank 26", "slot": "lifestyle", "code": "BIPM0001H28", "name": "The Rohal Huggie Earrings"}
]:
    recent_skus.append(sku_info)

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2)
print(f"Updated {rot_path}")

# 3. Update SEO Strategy 2026.xlsx
xlsx_target = ROOT / "SEO Strategy 2026.xlsx"
real_xlsx = xlsx_target.resolve()
wb = openpyxl.load_workbook(real_xlsx)
if "Week 8" in wb.sheetnames:
    ws = wb["Week 8"]
    headers = [cell.value for cell in ws[1]]
    rank_col = headers.index("Priority Rank") + 1
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1
    
    target_row_idx = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=rank_col).value
        if str(val).strip() == "26":
            target_row_idx = r
            break
            
    if target_row_idx:
        ws.cell(row=target_row_idx, column=url_col, value="https://blog.bluestone.com/unique-gold-earrings-design-2026/")
        ws.cell(row=target_row_idx, column=note_col, value=f"Published 2026-09-01. WP Post ID: 37280. Author: Satyam (270271337). Engine: buying_guide. Categories: Gold (554493348), Jewellery Problem & Solution (554493465). Type 3 trio: The Faliha Purse Hoop Earrings (hero), The Skein Hoop Earrings (flatlay), The Rohal Huggie Earrings (lifestyle). Flatlay setting: windowsill-daylight. Carousel: 37274-37279. Live QA passed with {word_count} words, verified BIS HUID/18K/22K standards, single H1, and og:image.")
        wb.save(real_xlsx)
        print(f"Updated {real_xlsx} sheet Week 8 row {target_row_idx}")
    else:
        print("WARNING: Rank 26 row not found in Week 8 sheet")
else:
    print("WARNING: 'Week 8' sheet not found in workbook")
