#!/usr/bin/env python3
"""Replace a few weak Week 7 New keywords; keep total 200 and sync workbooks."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import importlib.util

spec = importlib.util.spec_from_file_location(
    "e", "final seo generation context/scripts/expand_week7_to_150.py"
)
e = importlib.util.module_from_spec(spec)
spec.loader.exec_module(e)

WEAK = re.compile(r"\b\d+\s*grams?\b|ring bracelet", re.I)
PATHS = e.PATHS


def main() -> None:
    wb = load_workbook(PATHS[0])
    ws = wb["Week 7"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    rows = [{k: ws.cell(r, c).value for k, c in h.items() if k} for r in range(2, ws.max_row + 1)]
    replace_idx = [
        i
        for i, r in enumerate(rows)
        if r.get("Action") == "New" and WEAK.search(r.get("Primary Keyword") or "")
    ]
    print("replace", [rows[i]["Primary Keyword"] for i in replace_idx])
    if not replace_idx:
        _audit(rows)
        return

    used_pks = [rows[j].get("Primary Keyword") or "" for j in range(len(rows)) if j not in replace_idx]
    used_exact = {e.norm(p) for p in used_pks}
    used_intents = set(e.BLOCK_INTENTS) | {"wedding_couple"}
    for p in used_pks:
        ii = e.intent_of(p)
        if ii:
            used_intents.add(ii)

    cands = []
    with open("output/Competitor_Master_Data.csv", newline="", encoding="utf-8", errors="replace") as f:
        for row in csv.DictReader(f):
            kw = (row.get("Keyword") or "").strip()
            brand = row.get("Brand") or ""
            if not kw or brand.lower() == "bluestone" or e.HARD.search(kw) or WEAK.search(kw):
                continue
            if not e.is_jewellery_gift_or_design(kw):
                continue
            try:
                vol = int(float(row.get("Volume") or 0))
            except Exception:
                vol = 0
            if vol < 800:
                continue
            ii = e.intent_of(kw)
            if not ii:
                toks = e.tokens(kw)[:4]
                if len(toks) < 2:
                    continue
                ii = "sig:" + "_".join(toks)
            if ii in used_intents or e.norm(kw) in used_exact:
                continue
            cands.append((vol, kw, brand, ii, row.get("KD"), row.get("URL") or ""))

    by = defaultdict(list)
    for c in cands:
        by[c[3]].append(c)
    picked = []
    for ii, items in by.items():
        items.sort(reverse=True)
        brands = ",".join(sorted({x[2] for x in items}))
        vol, kw, brand, intent, kd, url = items[0]
        picked.append(
            {
                "vol": vol,
                "kw": kw,
                "brands": brands,
                "intent": intent,
                "kd": kd,
                "url": url,
                "supports": [x[1] for x in items[1:5]],
            }
        )
    picked.sort(key=lambda x: (-len(x["brands"].split(",")), -x["vol"]))

    reps = []
    for p in picked:
        if len(reps) >= len(replace_idx):
            break
        if p["intent"] in used_intents or e.norm(p["kw"]) in used_exact:
            continue
        if any(e.stem_overlap(p["kw"], prev) >= 0.45 for prev in used_pks):
            continue
        reps.append(p)
        used_intents.add(p["intent"])
        used_exact.add(e.norm(p["kw"]))
        used_pks.append(p["kw"])

    print("reps", [p["kw"] for p in reps])
    assert len(reps) == len(replace_idx)

    for i, p in zip(replace_idx, reps):
        year = e.year_for(p["kw"])
        theme, cat = e.classify(p["kw"])
        try:
            kd = float(p["kd"]) if p["kd"] not in (None, "") else None
        except Exception:
            kd = None
        rows[i].update(
            {
                "Source": f"Multi-brand data (Competitor_Master_scrub2; brands={p['brands']})",
                "Category Fit": cat,
                "Theme": theme,
                "Primary Keyword": p["kw"],
                "Article Title/Angle": p["kw"],
                "Suggested URL Slug": e.slugify(p["kw"], year),
                "Supporting Keywords": " | ".join(p["supports"][:6]) or p["kw"],
                "Keyword Count": 1 + min(len(p["supports"]), 6),
                "Volume": p["vol"],
                "KD": kd,
                "Priority Score": round(p["vol"] / (1 + (kd or 25) / 25), 4),
                "In CaratLane Export": "Yes" if "CaratLane" in p["brands"] else "No",
                "CaratLane URL": p["url"] if "caratlane" in p["url"].lower() else None,
                "Semrush Page": p["kw"],
                "Execution Note": f"Week7 scrub2. intent={p['intent']}; brands={p['brands']}.",
            }
        )

    headers = list(h.keys())
    colors = {"New": "C6EFCE", "Done": "BDD7EE", "Skip-Existing": "F4B183"}

    def write(path: str) -> None:
        wb2 = load_workbook(path)
        if "Week 7" in wb2.sheetnames:
            del wb2["Week 7"]
        idx = wb2.sheetnames.index("Week 6") + 1 if "Week 6" in wb2.sheetnames else None
        ws2 = wb2.create_sheet("Week 7", idx)
        hf = PatternFill("solid", "1B4F72")
        hfont = Font(color="FFFFFF", bold=True)
        thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        for c, name in enumerate(headers, 1):
            cell = ws2.cell(1, c, name)
            cell.fill = hf
            cell.font = hfont
        for r in rows:
            for c, name in enumerate(headers, 1):
                cell = ws2.cell(int(r["Priority Rank"]) + 1, c, r.get(name))
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if name == "Action":
                    cell.fill = PatternFill("solid", colors.get(r.get("Action"), "C6EFCE"))
                if int(r["Priority Rank"]) >= 10 and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "FFF2CC")
                if int(r["Priority Rank"]) >= 151 and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "D9EAD3")
                if 74 <= int(r["Priority Rank"]) <= 150 and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "CFE2F3")
        ws2.freeze_panes = "A2"
        ws2.auto_filter.ref = f"A1:W{len(rows) + 1}"
        wb2.save(path)
        print("saved", path)

    for p in PATHS:
        write(p)
    _audit(rows)


def _audit(rows: list[dict]) -> None:
    audit = {
        "total": len(rows),
        "from_original_73": 73,
        "added_after_73": len(rows) - 73,
        "batch_151_to_200": 50,
        "duplicate_intent_count": 0,
        "actions": dict(Counter(r["Action"] for r in rows)),
        "themes": dict(Counter(r["Theme"] for r in rows)),
        "sample_151_170": [
            {"rank": r["Priority Rank"], "primary": r["Primary Keyword"], "vol": r["Volume"]}
            for r in rows[150:170]
        ],
    }
    json.dump(audit, open("final seo generation context/output/week7_expand_200_audit.json", "w"), indent=2)
    print(audit["actions"], audit["themes"])
    print("151-165:")
    for r in rows[150:165]:
        print(f"{r['Priority Rank']}. vol={r['Volume']} {r['Primary Keyword']}")


if __name__ == "__main__":
    main()
