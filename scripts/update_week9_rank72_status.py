#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 72."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 72
PRIMARY = "wedding gold long necklace designs"
SLUG = "wedding-gold-long-necklace-designs-2026"
BLOG_URL = "https://blog.bluestone.com/wedding-gold-long-necklace-designs-2026/"
POST_ID = "40264"
CAROUSEL_MEDIA = "40258,40259,40260,40261,40262,40263"
TYPE3_MEDIA = "40265,40266,40267"
VISIBLE_WORDS = "3588"
NOTES = "Published 2026-09-26. WP#40264. Author Satyam 270271337. Type 3 trio: BVPJ0935C06 / BIAV0987N78 / BIMA1081C01. Flatlay: festive-mantel. 6 carousel cards live. BIS hallmarking HUID & net weight billing verified. Zero dash/price errors."

def update_status_csv():
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": TYPE3_MEDIA,
                        "lines": "",
                        "visible_words": VISIBLE_WORDS,
                        "notes": "2026-09-26T16:09:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": TYPE3_MEDIA,
            "lines": "",
            "visible_words": VISIBLE_WORDS,
            "notes": "2026-09-26T16:09:00Z | SERP intelligence pipeline v1"
        })

    # Sort by rank numeric
    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-26"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BVPJ0935C06 The Shubhlatika Mangalsutra Necklace",
        "flatlay": "BIAV0987N78 The Ailia Evil Eye Layered Necklace",
        "lifestyle": "BIMA1081C01 The Yeijah Mangaslsutra Necklace"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "festive-mantel",
        "note": "wedding gold long necklace designs The Ailia Evil Eye Layered Necklace on festive bridal mantel with brass diya, fresh mogra and marigold flowers, and raw silk runner"
    })

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"Updated {rotation_file} with Rank {RANK} rotation and flatlay setting")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print("Excel workbook not found!")
        return

    try:
        wb = openpyxl.load_workbook(xlsx_path)
        if "Week 9" not in wb.sheetnames:
            print("Sheet Week 9 not found!")
            return

        ws = wb["Week 9"]
        headers = [c.value for c in ws[1]]
        url_col = headers.index("Bluestone Blog URL") + 1
        note_col = headers.index("Execution Note") + 1
        rank_col = headers.index("Priority Rank") + 1

        updated_row = None
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=rank_col).value) == str(RANK):
                ws.cell(row=row, column=url_col, value=BLOG_URL)
                ws.cell(row=row, column=note_col, value=NOTES)
                updated_row = row
                break

        if updated_row:
            wb.save(xlsx_path)
            print(f"Updated Excel workbook row {updated_row} for Rank {RANK}")
        else:
            print(f"Row for Rank {RANK} not found in Excel workbook")
    except Exception as e:
        print(f"Warning: could not update Excel workbook: {e}")

if __name__ == "__main__":
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()
