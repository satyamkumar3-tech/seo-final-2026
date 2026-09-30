#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files and Excel workbook for Week 9 Rank 60: 916-hallmark-gold-2026."""
import os, sys, json, csv
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RANK = 60
PRIMARY = "916 hallmark gold"
SLUG = "916-hallmark-gold-2026"
POST_ID = 40141
URL = "https://blog.bluestone.com/916-hallmark-gold-2026/"
WORD_COUNT = "2644"
NOW_ISO = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S")

# 1. Update output/Week9_Blog_Queue_status.csv
status_csv_path = ROOT / "output" / "Week9_Blog_Queue_status.csv"
rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for r in reader:
            if r and str(r[0]).strip() != str(RANK):
                rows.append(r)

new_row = [str(RANK), "Done", PRIMARY, "Done", str(POST_ID), URL, WORD_COUNT, NOW_ISO]
rows.append(new_row)
rows.sort(key=lambda x: int(x[0]) if x[0].isdigit() else 9999)

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/Week9_Blog_Queue.csv
queue_csv_path = ROOT / "output" / "Week9_Blog_Queue.csv"
if queue_csv_path.exists():
    q_rows = []
    with open(queue_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        q_fields = reader.fieldnames
        for r in reader:
            if str(r.get("rank", "")).strip() == str(RANK):
                r["action"] = "Done"
            q_rows.append(r)
    with open(queue_csv_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=q_fields)
        writer.writeheader()
        writer.writerows(q_rows)
    print(f"Updated {queue_csv_path}")

# 3. Update output/product_rotation.json
rot_path = ROOT / "output" / "product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {}

recent_trios = rot_data.setdefault("recent_type3_trios", [])
recent_trios = [t for t in recent_trios if isinstance(t, dict) and t.get("rank") != f"Week9 Rank {RANK}" and t.get("rank") != RANK]
recent_trios.append({
    "rank": f"Week9 Rank {RANK}",
    "hero": "BISM0003O14 The Muricelle Bangle",
    "flatlay": "BISW1080P32 The Teshvarya Pendant",
    "lifestyle": "BIKR0993R117 The Luvee Highway Ring"
})
rot_data["recent_type3_trios"] = recent_trios

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings = [s for s in recent_settings if not isinstance(s, dict) or (s.get("rank") != f"Week9 Rank {RANK}" and s.get("rank") != RANK)]
recent_settings.append({
    "rank": f"Week9 Rank {RANK}",
    "setting": "desk-kraft",
    "note": "916 hallmark gold The Teshvarya Pendant on teak desk with brass loupe"
})
rot_data["recent_flatlay_settings"] = recent_settings

recent_skus = rot_data.setdefault("recent_skus", [])
recent_skus = [s for s in recent_skus if not isinstance(s, dict) or s.get("rank") != f"Week9 Rank {RANK}"]
recent_skus.extend([
    {"rank": f"Week9 Rank {RANK}", "slot": "hero", "code": "BISM0003O14", "name": "The Muricelle Bangle"},
    {"rank": f"Week9 Rank {RANK}", "slot": "flatlay", "code": "BISW1080P32", "name": "The Teshvarya Pendant"},
    {"rank": f"Week9 Rank {RANK}", "slot": "lifestyle", "code": "BIKR0993R117", "name": "The Luvee Highway Ring"}
])
rot_data["recent_skus"] = recent_skus
rot_data["updated"] = datetime.now().strftime("%Y-%m-%d")

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2, ensure_ascii=False)
print(f"Updated {rot_path}")

# 4. Update SEO Strategy 2026.xlsx Sheet Week 9
xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
if xlsx_path.exists():
    try:
        wb = openpyxl.load_workbook(xlsx_path)
        if "Week 9" in wb.sheetnames:
            ws = wb["Week 9"]
            headers = [cell.value for cell in ws[1]]
            rank_col = headers.index("Priority Rank") + 1 if "Priority Rank" in headers else 1
            url_col = headers.index("Bluestone Blog URL") + 1 if "Bluestone Blog URL" in headers else None
            note_col = headers.index("Execution Note") + 1 if "Execution Note" in headers else None
            action_col = headers.index("Action") + 1 if "Action" in headers else None

            row_found = False
            for r in range(2, ws.max_row + 1):
                val = ws.cell(row=r, column=rank_col).value
                if val == RANK or str(val) == str(RANK):
                    if url_col:
                        ws.cell(row=r, column=url_col, value=URL)
                    if action_col:
                        ws.cell(row=r, column=action_col, value="Done")
                    if note_col:
                        old_note = ws.cell(row=r, column=note_col).value or ""
                        new_note = f"Published 2026-09-25 as WP post ID {POST_ID}. Engine: buying_guide; {WORD_COUNT} visible words; 10 H2s; 3 BIS hallmark marks; 750 gold comparison; 3D coverflow carousel (40135-40140); Type 3 media (40142-40144); flatlay setting: desk-kraft; Author: Satyam (270271337); passed live QA."
                        ws.cell(row=r, column=note_col, value=new_note)
                    row_found = True
                    print(f"Updated Excel row for Rank {RANK} in Sheet 'Week 9'")
                    break
            if row_found:
                wb.save(xlsx_path)
                print(f"Saved changes to {xlsx_path}")
    except Exception as e:
        print(f"Excel update warning: {e}")
