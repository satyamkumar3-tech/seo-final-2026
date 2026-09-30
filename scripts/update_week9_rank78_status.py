#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 78: Indian Gold Earrings Designs Hoops."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 78
PRIMARY = "indian gold earrings designs hoops"
SLUG = "indian-gold-earrings-designs-hoops-2026"
BLOG_URL = "https://blog.bluestone.com/indian-gold-earrings-designs-hoops-2026/"
POST_ID = "40336"
CAROUSEL_MEDIA = "40330,40331,40332,40333,40334,40335"
TYPE3_MEDIA = "40337,40338,40339"
VISIBLE_WORDS = "2961"
NOTES = "Published 2026-09-26. WP#40336. Author Satyam 270271337. Type 3 trio: BIIP0427H16 (The Vicky Hoop Earrings) / BIJP0686H03 (The Faliha Purse Hoop Earrings) / BISP0427H21 (The Ursa Hoop Earrings). Flatlay: windowsill-daylight. 6 carousel cards live. BIS hallmarking HUID & net gold weight billing verified. Zero dash/price errors."

def update_status_csv():
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": TYPE3_MEDIA,
                        "lines": "",
                        "visible_words": VISIBLE_WORDS,
                        "notes": "2026-09-26T22:45:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": TYPE3_MEDIA,
            "lines": "",
            "visible_words": VISIBLE_WORDS,
            "notes": "2026-09-26T22:45:00Z | SERP intelligence pipeline v1"
        })

    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-26"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BIIP0427H16 The Vicky Hoop Earrings",
        "flatlay": "BIJP0686H03 The Faliha Purse Hoop Earrings",
        "lifestyle": "BISP0427H21 The Ursa Hoop Earrings"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "windowsill-daylight",
        "note": f"{PRIMARY} The Faliha Purse Hoop Earrings in ceramic dish on oak windowsill with brass loupe"
    })

    ranks = data.setdefault("ranks", {})
    ranks[str(RANK)] = {
        "primary": PRIMARY,
        "slug": SLUG,
        "type3_skus": ["BIIP0427H16", "BIJP0686H03", "BISP0427H21"],
        "flatlay_setting": "windowsill-daylight",
        "carousel_skus": [
            "BIPM0001H28",
            "BISA0255D05",
            "BISP0427H21",
            "BINK0363H03",
            "BIIP0279S08",
            "BIPN0880H218"
        ]
    }

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    print(f"Updated {rotation_file} with Rank {RANK} rotation details")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print(f"Warning: {xlsx_path} not found")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    sheet = wb["Week 9"]

    headers = [cell.value for cell in sheet[1]]
    url_col = headers.index("Bluestone Blog URL") + 1 if "Bluestone Blog URL" in headers else None
    slug_col = headers.index("Suggested URL Slug") + 1 if "Suggested URL Slug" in headers else None
    note_col = headers.index("Execution Note") + 1 if "Execution Note" in headers else None
    action_col = headers.index("Action") + 1 if "Action" in headers else None

    row_found = False
    for r in range(2, sheet.max_row + 1):
        cell_val = sheet.cell(row=r, column=1).value
        if str(cell_val) == str(RANK):
            if url_col:
                sheet.cell(row=r, column=url_col, value=BLOG_URL)
            if slug_col:
                sheet.cell(row=r, column=slug_col, value=SLUG)
            if note_col:
                sheet.cell(row=r, column=note_col, value=NOTES)
            if action_col:
                sheet.cell(row=r, column=action_col, value="Done")
            row_found = True
            print(f"Updated XLSX row {r} for Rank {RANK}")
            break

    if row_found:
        wb.save(xlsx_path)
        print(f"Saved changes to {xlsx_path}")
    else:
        print(f"Could not find Rank {RANK} row in {xlsx_path}")

if __name__ == "__main__":
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()
