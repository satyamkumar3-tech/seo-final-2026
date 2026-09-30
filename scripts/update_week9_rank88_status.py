#!/usr/bin/env python3
"""
Update status files and SEO strategy spreadsheet for Week 9 Rank 88.
"""
import csv
import json
import os
import openpyxl

def update_csv():
    csv_path = "output/Week9_Blog_Queue_status.csv"
    rows = []
    headers = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        headers = next(reader)
        for row in reader:
            if row:
                rows.append(row)
    
    # Check if rank 88 already exists
    rank_88_row = [
        "88",
        "panna stone ring",
        "panna-stone-ring-2026",
        "https://blog.bluestone.com/panna-stone-ring-2026/",
        "40451",
        "Done",
        "40447,40448,40449,40450,37070,37073",
        "40452,40453,40454",
        "",
        "4493",
        "2026-09-28T17:15:00+05:30 | SERP intelligence pipeline v1"
    ]
    
    found = False
    for i, row in enumerate(rows):
        if row and row[0] == "88":
            rows[i] = rank_88_row
            found = True
            break
    if not found:
        rows.append(rank_88_row)
        
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    print(f"Updated {csv_path} with Rank 88 status (found={found}).")

def update_json():
    json_path = "output/product_rotation.json"
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    data["updated"] = "2026-09-28"
    
    # recent_type3_trios
    trios = data.get("recent_type3_trios", [])
    # Check if rank 88 already in trios
    trios = [t for t in trios if str(t.get("rank")) not in ("88", "Week9 Rank 88")]
    trios.append({
        "rank": "Week9 Rank 88",
        "hero": "BISL0851R28 The Jasper Band For Him",
        "flatlay": "BIAR0097R16 The Quinn Ring",
        "lifestyle": "BINS0639R18 The Gigi Ring"
    })
    data["recent_type3_trios"] = trios
    
    # recent_skus
    recent_skus = data.get("recent_skus", [])
    recent_skus = [s for s in recent_skus if str(s.get("rank")) not in ("88", "Week9 Rank 88")]
    recent_skus.extend([
        {"rank": "Week9 Rank 88", "slot": "hero", "code": "BISL0851R28", "name": "The Jasper Band For Him"},
        {"rank": "Week9 Rank 88", "slot": "flatlay", "code": "BIAR0097R16", "name": "The Quinn Ring"},
        {"rank": "Week9 Rank 88", "slot": "lifestyle", "code": "BINS0639R18", "name": "The Gigi Ring"}
    ])
    data["recent_skus"] = recent_skus
    
    # recent_flatlay_settings
    settings = data.get("recent_flatlay_settings", [])
    settings.append("cafe-tray")
    if len(settings) > 6:
        settings = settings[-6:]
    data["recent_flatlay_settings"] = settings
    
    # ranks
    if "ranks" not in data:
        data["ranks"] = {}
    data["ranks"]["88"] = {
        "rank": "88",
        "primary": "panna stone ring",
        "slug": "panna-stone-ring-2026",
        "hero_sku": "BISL0851R28",
        "flatlay_sku": "BIAR0097R16",
        "lifestyle_sku": "BINS0639R18",
        "flatlay_setting": "cafe-tray"
    }
    
    # posts
    posts = data.get("posts", [])
    posts = [p for p in posts if str(p.get("rank")) not in ("88", "Week9 Rank 88")]
    posts.append({
        "rank": "Week9 Rank 88",
        "slug": "panna-stone-ring-2026",
        "flatlay_setting": "cafe-tray",
        "skus": {
            "hero": "BISL0851R28",
            "flatlay": "BIAR0097R16",
            "lifestyle": "BINS0639R18"
        },
        "names": {
            "hero": "The Jasper Band For Him",
            "flatlay": "The Quinn Ring",
            "lifestyle": "The Gigi Ring"
        },
        "carousel_skus": ["BISV0910R24", "BISL0851R28", "BIAR0097R04", "BIIP0090R24", "BISE0932R181", "BIAR0097R16"]
    })
    data["posts"] = posts
    
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print(f"Updated {json_path} with Rank 88 rotation data.")

def update_excel():
    excel_path = "SEO Strategy 2026.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Week 9"]
    
    # Find row with Priority Rank 88
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == "88":
            target_row = r
            break
            
    if not target_row:
        raise ValueError("Could not find Priority Rank 88 in Sheet 'Week 9'!")
        
    print(f"Found Priority Rank 88 at row {target_row}")
    # Column 12: Bluestone Blog URL
    # Column 23: Execution Note
    ws.cell(target_row, 12).value = "https://blog.bluestone.com/panna-stone-ring-2026/"
    ws.cell(target_row, 23).value = "Published live post 40451. Verified live QA 2026-09-28."
    
    wb.save(excel_path)
    print(f"Updated {excel_path} Sheet 'Week 9' row {target_row}.")

if __name__ == "__main__":
    update_csv()
    update_json()
    update_excel()
    print("All status updates completed successfully!")
