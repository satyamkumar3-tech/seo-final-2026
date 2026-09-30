#!/usr/bin/env python3
"""Update status files and Excel workbook for Week 8 Rank 1: Diamond Rings."""
import os
import csv
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

# 1. Update output/Week8_Blog_Queue_status.csv
status_csv_path = ROOT / "output/Week8_Blog_Queue_status.csv"
fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            if r["rank"] != "1":
                rows.append(r)

rank1_row = {
    "rank": "1",
    "primary": "diamond rings",
    "slug": "diamond-rings-2026",
    "blog_url": "https://blog.bluestone.com/diamond-rings-2026/",
    "wp_post_id": "36937",
    "status": "published",
    "carousel_media": "36931-36936",
    "type3_media": "hero 36938; flatlay 36939; lifestyle 36940",
    "lines": "115",
    "visible_words": "3007",
    "notes": "Published 2026-08-31 via buying_guide engine. 4Cs diamond education, 3-diamond trilogy symbolism, 14K/18K gold settings, BIS HUID hallmarking, SGL/IGI/GIA certification, 6-card carousel, 3 Type 3 images (The Rafia Ring, The Malibu Ring, The Ebony Ring), passed live QA."
}
rows.append(rank1_row)
rows.sort(key=lambda x: int(x["rank"]))

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/product_rotation.json
rot_path = ROOT / "output/product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {"recent_flatlay_settings": [], "recent_skus": []}

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings.append({
    "rank": "Week8 Rank 1",
    "setting": "marble-vanity",
    "note": "diamond rings The Malibu Ring"
})

recent_skus = rot_data.setdefault("recent_skus", [])
for sku_info in [
    {"rank": "Week8 Rank 1", "slot": "hero", "code": "BIAB0503R03", "name": "The Rafia Ring"},
    {"rank": "Week8 Rank 1", "slot": "flatlay", "code": "BIPM0017R18", "name": "The Malibu Ring"},
    {"rank": "Week8 Rank 1", "slot": "lifestyle", "code": "BIIP0090R24", "name": "The Ebony Ring"}
]:
    recent_skus.append(sku_info)

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2)
print(f"Updated {rot_path}")

# 3. Update SEO Strategy 2026.xlsx
xlsx_target = ROOT / "SEO Strategy 2026.xlsx"
# Check if symlink or regular file
real_xlsx = xlsx_target.resolve()
wb = openpyxl.load_workbook(real_xlsx)
if "Week 8" in wb.sheetnames:
    ws = wb["Week 8"]
    headers = [cell.value for cell in ws[1]]
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1
    
    ws.cell(row=2, column=url_col, value="https://blog.bluestone.com/diamond-rings-2026/")
    ws.cell(row=2, column=note_col, value="Published 2026-08-31. WP Post ID: 36937. Author: Satyam (270271337). Engine: buying_guide. Categories: Diamond (554493418), Ring (554493326), diamond ring (554493475), Jewellery Problem & Solution (554493465). Type 3 trio: The Rafia Ring (hero 36938), The Malibu Ring (flatlay 36939), The Ebony Ring (lifestyle 36940). Flatlay setting: marble-vanity. Carousel: 36931-36936. Live QA passed with 3007 words, verified 4Cs/HUID/SGL/IGI/GIA standards, single H1, and og:image.")
    wb.save(real_xlsx)
    print(f"Updated {real_xlsx} sheet Week 8 row 2")
else:
    print("WARNING: 'Week 8' sheet not found in workbook")

