#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files and SEO strategy spreadsheet for Week 9 Rank 89."""
import csv
import json
import os
import openpyxl

def update_csv():
    csv_path = "output/Week9_Blog_Queue_status.csv"
    rows = []
    headers = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        headers = next(reader)
        for row in reader:
            if row:
                rows.append(row)
    
    rank_89_row = [
        "89",
        "daily wear gold earrings for women",
        "daily-wear-gold-earrings-for-women-2026",
        "https://blog.bluestone.com/daily-wear-gold-earrings-for-women-2026/",
        "40462",
        "Done",
        "40456,40457,40458,40459,40460,40461",
        "40463,40464,40465",
        "",
        "2455",
        "2026-09-28T18:04:00+05:30 | SERP intelligence pipeline v1"
    ]
    
    found = False
    for i, row in enumerate(rows):
        if row and row[0] == "89":
            rows[i] = rank_89_row
            found = True
            break
    if not found:
        rows.append(rank_89_row)
        
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    print(f"Updated {csv_path} with Rank 89 status (found={found}).")

def update_json():
    json_path = "output/product_rotation.json"
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    data["updated"] = "2026-09-28"
    
    # recent_type3_trios
    trios = data.get("recent_type3_trios", [])
    trios = [t for t in trios if str(t.get("rank")) not in ("89", "Week9 Rank 89")]
    trios.append({
        "rank": "Week9 Rank 89",
        "hero": "BIPM0001H28 The Rohal Huggie Earrings",
        "flatlay": "BIPN0880H218 The Nettile Huggie Earrings",
        "lifestyle": "BIIP0427H16 The Vicky Hoop Earrings"
    })
    data["recent_type3_trios"] = trios
    
    # recent_skus
    recent_skus = data.get("recent_skus", [])
    recent_skus = [s for s in recent_skus if str(s.get("rank")) not in ("89", "Week9 Rank 89")]
    recent_skus.extend([
        {"rank": "Week9 Rank 89", "slot": "hero", "code": "BIPM0001H28", "name": "The Rohal Huggie Earrings"},
        {"rank": "Week9 Rank 89", "slot": "flatlay", "code": "BIPN0880H218", "name": "The Nettile Huggie Earrings"},
        {"rank": "Week9 Rank 89", "slot": "lifestyle", "code": "BIIP0427H16", "name": "The Vicky Hoop Earrings"}
    ])
    data["recent_skus"] = recent_skus
    
    # recent_flatlay_settings
    settings = data.get("recent_flatlay_settings", [])
    settings.append("marble-vanity")
    if len(settings) > 6:
        settings = settings[-6:]
    data["recent_flatlay_settings"] = settings
    
    # ranks
    if "ranks" not in data:
        data["ranks"] = {}
    data["ranks"]["89"] = {
        "rank": "89",
        "primary": "daily wear gold earrings for women",
        "slug": "daily-wear-gold-earrings-for-women-2026",
        "hero_sku": "BIPM0001H28",
        "flatlay_sku": "BIPN0880H218",
        "lifestyle_sku": "BIIP0427H16",
        "flatlay_setting": "marble-vanity"
    }
    
    # posts
    posts = data.get("posts", [])
    posts = [p for p in posts if str(p.get("rank")) not in ("89", "Week9 Rank 89")]
    posts.append({
        "rank": "Week9 Rank 89",
        "slug": "daily-wear-gold-earrings-for-women-2026",
        "flatlay_setting": "marble-vanity",
        "skus": {
            "hero": "BIPM0001H28",
            "flatlay": "BIPN0880H218",
            "lifestyle": "BIIP0427H16"
        },
        "names": {
            "hero": "The Rohal Huggie Earrings",
            "flatlay": "The Nettile Huggie Earrings",
            "lifestyle": "The Vicky Hoop Earrings"
        },
        "carousel_skus": ["BISA0255D05", "BINK0363H03", "BIIP0279S08", "BIJP0686H03", "BISP0427H21", "BIPM0001H28"]
    })
    data["posts"] = posts
    
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print(f"Updated {json_path} with Rank 89 rotation data.")

def update_excel():
    excel_path = "SEO Strategy 2026.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Week 9"]
    
    target_row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == "89":
            target_row = r
            break
            
    if not target_row:
        raise ValueError("Could not find Priority Rank 89 in Sheet 'Week 9'!")
        
    print(f"Found Priority Rank 89 at row {target_row}")
    # Column 12: Bluestone Blog URL
    # Column 23: Execution Note
    ws.cell(target_row, 12).value = "https://blog.bluestone.com/daily-wear-gold-earrings-for-women-2026/"
    ws.cell(target_row, 23).value = "Published live post 40462. Verified live QA 2026-09-28."
    
    wb.save(excel_path)
    print(f"Updated {excel_path} Sheet 'Week 9' row {target_row}.")

if __name__ == "__main__":
    update_csv()
    update_json()
    update_excel()
    print("All status updates completed successfully!")
