#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files and Excel workbook for Week 9 Rank 65."""
import os
import json
import csv
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

RANK = 65
PRIMARY_KW = "blue stone ring for men"
POST_ID = 40188
LIVE_URL = "https://blog.bluestone.com/blue-stone-ring-for-men-2026/"
WORD_COUNT = 2419
TIMESTAMP = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S")

# 1. Upsert Week9_Blog_Queue_status.csv
status_file = ROOT / "output/Week9_Blog_Queue_status.csv"
rows = []
header_present = False
if status_file.exists():
    with open(status_file, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r in reader:
            if not r:
                continue
            if r[0] == "rank":
                header_present = True
                rows.append(r)
            elif r[0] == str(RANK):
                continue  # replace
            else:
                rows.append(r)

new_row = [str(RANK), "Done", PRIMARY_KW, "Done", str(POST_ID), LIVE_URL, str(WORD_COUNT), TIMESTAMP]
rows.append(new_row)

def rank_key(row):
    try:
        return int(row[0])
    except ValueError:
        return -1

if header_present:
    sorted_rows = [rows[0]] + sorted(rows[1:], key=rank_key)
else:
    sorted_rows = sorted(rows, key=rank_key)

with open(status_file, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(sorted_rows)
print(f"Upserted Rank {RANK} in {status_file}")

# 2. Upsert product_rotation.json
rot_file = ROOT / "output/product_rotation.json"
with open(rot_file, "r", encoding="utf-8") as f:
    rot_data = json.load(f)

rot_data["updated"] = datetime.now(timezone.utc).strftime("%Y-%m-%d")

trios = rot_data.get("recent_type3_trios", [])
trios = [t for t in trios if isinstance(t, dict) and str(t.get("rank")) not in [str(RANK), f"Week9 Rank {RANK}"]]
trios.append({
    "rank": f"Week9 Rank {RANK}",
    "hero": "BISV0910R24 The Interlink Band Ring",
    "flatlay": "BIIP0090R24 The Ebony Ring",
    "lifestyle": "BISL0851R28 The Jasper Band For Him"
})
rot_data["recent_type3_trios"] = trios

settings = rot_data.get("recent_flatlay_settings", [])
clean_settings = []
for s in settings:
    if isinstance(s, dict):
        if str(s.get("rank")) not in [str(RANK), f"Week9 Rank {RANK}"]:
            clean_settings.append(s)
    else:
        clean_settings.append(s)

clean_settings.append({
    "rank": f"Week9 Rank {RANK}",
    "setting": "cafe-tray",
    "note": "blue stone ring The Ebony Ring on stoneware cafe tray with espresso cup and brass loupe"
})
rot_data["recent_flatlay_settings"] = clean_settings

with open(rot_file, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2, ensure_ascii=False)
print(f"Upserted Rank {RANK} in {rot_file}")

# 3. Update SEO Strategy 2026.xlsx
xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
wb = openpyxl.load_workbook(xlsx_path)
ws = wb["Week 9"]

headers = [cell.value for cell in ws[1]]
rank_col = headers.index("Priority Rank") + 1
url_col = headers.index("Bluestone Blog URL") + 1
exec_col = headers.index("Execution Note") + 1

found = False
for r in range(2, ws.max_row + 1):
    val = ws.cell(row=r, column=rank_col).value
    if val == RANK or str(val).strip() == str(RANK):
        ws.cell(row=r, column=url_col, value=LIVE_URL)
        exec_note = f"Published 2026-09-25. WP#{POST_ID}. Author Satyam 270271337. Type 3 trio: BISV0910R24 / BIIP0090R24 / BISL0851R28. Flatlay: cafe-tray. 6 carousel cards live. BIS hallmarking & gemstone guidance verified. Zero dash/price errors."
        ws.cell(row=r, column=exec_col, value=exec_note)
        found = True
        print(f"Updated row {r} in Excel sheet Week 9")
        break

if found:
    wb.save(xlsx_path)
    print(f"Saved {xlsx_path}")
else:
    print(f"Warning: Rank {RANK} not found in Excel sheet Week 9")
