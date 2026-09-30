#!/usr/bin/env python3
"""Expand Week 7 to TARGET intent-unique rows using multi-brand competitor data."""

from __future__ import annotations

import csv
import json
import os
import re
from collections import Counter, defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PATHS = [
    "SEO Strategy 2026.xlsx",
    "final seo generation context/SEO Strategy 2026.xlsx",
    "article generation seo codex/SEO Strategy 2026.xlsx",
]
SRC = PATHS[0]
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TARGET = 200
FREEZE_THROUGH = 150  # keep existing ranks 1..FREEZE_THROUGH; fill above


def repo_path(*parts: str) -> str:
    return os.path.join(ROOT, *parts)


def norm(s: str) -> str:
    s = (s or "").strip().lower()
    s = s.replace("marriage", "wedding").replace("newly married", "newlywed").replace("new married", "newlywed")
    s = re.sub(r"[''`]", "", s)
    s = re.sub(r"[^a-z0-9\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


STOP = set(
    "a an the for to of in on with from and or my your our very best top unique useful "
    "luxury indian special new ideas idea suggestions suggestion gift gifts how design "
    "designs in gold silver women mens men man girl girls ladies under".split()
)


def tokens(s: str) -> list[str]:
    out = []
    for t in norm(s).split():
        if not t or t in STOP or re.fullmatch(r"20\d{2}|\d+", t):
            continue
        # light plural stem for overlap (bands->band, earrings->earring)
        if len(t) > 4 and t.endswith("s") and not t.endswith("ss"):
            t = t[:-1]
        out.append(t)
    return out


def stem_overlap(a: str, b: str) -> float:
    sa, sb = set(tokens(a)), set(tokens(b))
    if not sa or not sb:
        return 0.0
    return len(sa & sb) / len(sa | sb)


RULES = [
    ("karwa_wife", r"karwa|karva"),
    ("rakhi_sister", r"(rakhi|raksha).*(sister)|sister.*(rakhi|raksha)"),
    ("rakhi_brother", r"(rakhi|raksha).*(brother)|brother.*(rakhi|raksha)"),
    ("dhanteras", r"dhanteras|dhanatrayodashi"),
    ("wedding_bride", r"\b(bride|bridal|bride to be)\b"),
    ("wedding_friend", r"wedding.*friend|friend.*wedding"),
    ("wedding_girls_her", r"wedding gift.*(girl|women)|gift for married girl"),
    (
        "wedding_couple",
        r"(wedding|newlywed).*(couple)|gift for wedding|gifts for wedding|best wedding gift|marriage gift|reception gift|gift ideas for married",
    ),
    ("engagement_gift", r"engagement gift|gift.*engagement|propose day gift"),
    ("anniversary_parents", r"anniversary.*(parent|mom|dad)|mom dad anniversary"),
    ("anniversary_wife", r"anniversary.*wife"),
    ("anniversary_husband", r"anniversary.*(husband|hubby|him)\b"),
    ("sister_birthday", r"birthday.*sister|sister birthday"),
    ("brother_birthday", r"birthday.*brother|brother birthday"),
    ("mom_birthday", r"birthday.*(mom|mother|mummy)"),
    ("dad_birthday", r"birthday.*(dad|father)"),
    ("wife_birthday", r"birthday.*wife"),
    ("husband_birthday", r"birthday.*husband"),
    ("gf_birthday", r"birthday.*(girlfriend|gf)"),
    ("bf_birthday", r"birthday.*(boyfriend|bf)"),
    ("mothers_day_gift", r"mothers? day|mom'?s day"),
    ("fathers_day_gift", r"fathers? day"),
    ("valentine_husband", r"valentine.*husband"),
    ("valentine_wife", r"valentine.*wife"),
    ("valentine_girlfriend", r"valentine.*(girlfriend|gf)"),
    ("valentine_boyfriend", r"valentine.*(boyfriend|bf)"),
    ("secret_santa", r"secret santa"),
    ("christmas_gift", r"christmas|xmas"),
    ("diwali_gift", r"diwali gift|diwali.*gift|gift.*diwali"),
    ("new_year_gift", r"new year gift"),
    ("friendship_gift", r"friendship"),
    ("new_mom", r"new mom|new mother|baby shower"),
    ("first_night_wife", r"first night"),
    ("housewarming_gift", r"house.?warming|griha pravesh"),
    ("gift_sil", r"sister in law|bhabhi"),
    ("gift_mil", r"mother in law"),
    ("gift_fil", r"father in law"),
    ("gift_bil", r"brother in law"),
    ("gift_sister_general", r"gift for sister|best gift for sister|gift ideas for sister"),
    ("gift_mom_general", r"^gift for mom$|^gifts for mother$|^gift for mummy$|^mum gifts$|^gifts for parents$"),
    ("gift_dad_general", r"^gifts for father$|^gift for dad$"),
    ("gift_wife_general", r"^gift for wife$|^gifts for wife$|best gift for wife"),
    ("gift_husband_general", r"^gift for husband$|^gifts for husband$|best gift for husband"),
    ("akshaya_gift", r"akshaya"),
    ("gold_purity", r"purity|identify gold|hallmark|gold check|pure gold not suitable|how to check gold"),
    ("bangle_size", r"bangle size"),
    ("ring_size", r"ring size|finger ring size"),
    ("gst_gold", r"\bgst\b|gold tax"),
    ("buy_gold_day", r"best day to buy gold|when to buy gold"),
    ("making_charges", r"making charges"),
    ("karat_explain", r"\b(9|14|18|22)\s*(k|kt|karat|carat)\b"),
    ("demi_fine", r"demi fine"),
    ("name_ring", r"name ring"),
    ("upper_ear", r"upper ear|upper lobe|cartilage|ear piercing"),
    ("ear_saree", r"earrings for saree|earrings for gown"),
    ("ganesh_pendant", r"ganesh|ganpati pendant"),
    ("heart_pendant", r"heart.*(pendant|necklace)"),
    ("bangles_design", r"bangles design|latest.*bangles|gold bangle design"),
    ("10g_necklace", r"10 gram"),
    ("cocktail_ring", r"cocktail ring"),
    ("single_stone", r"single stone|solitaire"),
    ("lotus_jewellery", r"lotus"),
    ("hypoallergenic", r"hypoallergenic"),
    ("nose_pin", r"nose pin|nose piercing|nath"),
    ("kids_earrings", r"kids.*earring|earring.*kids"),
    ("kids_bracelet", r"kids.*(bracelet|bangle)|baby bracelet|kids bangles"),
    ("mens_chain", r"(men|mens).*chain|chain for men|heavy gold chain"),
    ("mens_bracelet", r"(men|mens).*bracelet|bracelet for men"),
    ("mens_ring", r"(men|mens).*\bring\b|\bring\b for men|gents ring|black ring for men"),
    ("mens_hoop", r"(men|mens).*hoop|hoop.*\bmen"),
    ("teen_bracelet", r"teenage"),
    ("necklace_types", r"types of necklaces"),
    ("mangalsutra_design", r"mangalsutra"),
    ("daily_wear_earrings", r"daily wear.*earring|office wear.*earring"),
    ("daily_wear_jewellery", r"daily wear.*(jewellery|jewelry|gold)"),
    ("choker", r"choker"),
    ("pearl_earrings", r"pearl earring"),
    ("kundan", r"kundan"),
    ("temple_jewellery", r"temple (jewellery|jewelry|set|necklace|earrings)"),
    ("oxidised_jewellery", r"oxidis"),
    ("gold_coin_gift", r"gold coin"),
    ("jhumka", r"jhumka|jhumki"),
    ("hoop_earrings", r"hoop earring"),
    ("stud_earrings", r"stud earring"),
    ("layer_necklace", r"layer.*necklace|multi layer|double layer|three layer"),
    ("tennis_bracelet", r"tennis bracelet"),
    ("charm_bracelet", r"charm bracelet"),
    ("evil_eye_bracelet", r"evil eye.*(bracelet|pendant)"),
    ("evil_eye_anklet", r"evil eye anklet"),
    ("initial_jewellery", r"initial (necklace|pendant|ring)|alphabet charm"),
    ("birthstone", r"birthstone"),
    ("toe_ring", r"toe ring|bichiya"),
    ("anklet", r"anklet|payal"),
    ("maang_tikka", r"maang tikka|matha patti|mathapatti"),
    ("waist_chain", r"waist chain|kamarband"),
    ("hathphool", r"hathphool|haath phool"),
    ("couple_rings", r"couple ring"),
    ("stackable_rings", r"stackable|stack ring"),
    ("bracelet_for_women", r"bracelet for women|womens bracelet|gold bracelet for women"),
    ("earrings_for_wedding", r"earrings for wedding|wedding earrings"),
    ("meenakari", r"meenakari|meena work"),
    ("polki", r"\bpolki\b"),
    ("jadau", r"\bjadau\b"),
    ("antique_jewellery", r"antique (jewellery|jewelry)"),
    ("rose_gold_jewellery", r"rose gold"),
    ("white_gold_jewellery", r"white gold"),
    ("platinum_jewellery", r"platinum (ring|band|bands|jewellery|jewelry|chain|bracelet)"),
    ("diamond_earrings", r"diamond earring"),
    ("diamond_necklace", r"diamond necklace"),
    ("diamond_bracelet", r"diamond bracelet"),
    ("diamond_pendant", r"diamond pendant"),
    ("gold_necklace_set", r"necklace set|necklace earring set"),
    ("lightweight_jewellery", r"lightweight|light weight"),
    ("office_wear_jewellery", r"office wear"),
    ("party_wear_jewellery", r"party wear"),
    ("traditional_jewellery", r"traditional (jewellery|jewelry)"),
    ("modern_jewellery", r"modern (jewellery|jewelry)"),
    ("minimalist_jewellery", r"minimalist|minimal jewellery"),
    ("statement_necklace", r"statement necklace"),
    ("bajuband", r"bajuband|armlet"),
    ("gold_chain_women", r"gold chain for women|womens gold chain"),
    ("box_chain", r"box chain"),
    ("rope_chain", r"rope chain"),
    ("figaro_chain", r"figaro"),
    ("cuba_chain", r"cuba chain|cuban chain"),
    ("pendant_for_men", r"pendant for men|mens pendant"),
    ("kada_men", r"\bkada\b.*men|men.*\bkada\b|mens kada"),
    ("kada_women", r"\bkada\b"),
    ("gifting_jewellery", r"jewellery gift|jewelry gift|gift jewellery"),
    # Extra pillars for 150→200 fill
    ("gold_tops", r"ear tops?|gold tops"),
    ("bali_earrings", r"\bbali\b"),
    ("chandbali", r"chandbali|chand bali"),
    ("drop_earrings", r"drop earring|dangler|dangle earring"),
    ("cluster_earrings", r"cluster earring"),
    ("huggie_earrings", r"huggie"),
    ("cartilage_earrings", r"cartilage|helix|tragus"),
    ("nose_ring", r"nose ring|nathni"),
    ("gold_kangan", r"\bkangan\b"),
    ("single_bangle", r"single bangle|one bangle"),
    ("pair_bangles", r"pair of bangles|bangle pair"),
    ("gold_set", r"gold set|jewellery set|jewelry set|necklace set"),
    ("mangalsutra_bracelet", r"mangalsutra bracelet"),
    ("black_beads", r"black beads|nallapusalu"),
    ("rudraksha_jewellery", r"rudraksha"),
    ("om_pendant", r"\bom\b.*(pendant|locket)|om pendant"),
    ("locket", r"\blocket\b"),
    ("charm_necklace", r"charm necklace"),
    ("name_necklace", r"name necklace|name pendant"),
    ("butterfly_jewellery", r"butterfly"),
    ("floral_jewellery", r"floral|flower (pendant|earring|ring|necklace)"),
    ("heart_jewellery", r"heart (ring|earring|bracelet)"),
    ("infinity_jewellery", r"infinity"),
    ("zodiac_jewellery", r"zodiac|rasi"),
    ("navaratna", r"navaratna|navratan"),
    ("panchaloha", r"panchaloha"),
    ("temple_earrings", r"temple earring"),
    ("south_indian_jewellery", r"south indian|tamil|kerala jewellery|andhra"),
    ("gujarati_jewellery", r"gujarati"),
    ("punjabi_jewellery", r"punjabi|jodha"),
    ("bridal_set", r"bridal set|bridal jewellery set"),
    ("engagement_ring_women", r"engagement ring"),
    ("promise_ring", r"promise ring"),
    ("cocktail_earrings", r"cocktail earring"),
    ("party_earrings", r"party wear earring|party earring"),
    ("office_earrings", r"office wear earring|office earring"),
    ("daily_necklace", r"daily wear necklace|daily necklace"),
    ("daily_bangles", r"daily wear bangle|daily bangle"),
    ("short_necklace", r"short necklace|short chain"),
    ("long_necklace", r"long necklace|long chain"),
    ("choker_set", r"choker set"),
    ("hasli", r"\bhasli\b"),
    ("haar", r"\bhaar\b|long haar"),
    ("mangalsutra_chain", r"mangalsutra chain"),
    ("thali_chain", r"thali|thaali"),
    ("gold_coin_pendant", r"coin pendant|gold coin pendant"),
    ("silver_anklet", r"silver anklet|silver payal"),
    ("gold_anklet", r"gold anklet|gold payal"),
    ("toe_ring_silver", r"silver toe|bichiya"),
    ("kids_necklace", r"kids.*(necklace|chain|pendant)|baby necklace"),
    ("kids_ring", r"kids ring|baby ring"),
    ("mens_kada_silver", r"silver kada"),
    ("mens_bracelet_leather", r"leather bracelet"),
    ("beaded_bracelet", r"beaded bracelet|bead bracelet"),
    ("cuff_bracelet", r"cuff bracelet|cuff bangle"),
    ("bangle_bracelet", r"bangle bracelet"),
    ("watch_jewellery", r"watch (bracelet|jewellery|jewelry)"),
    ("brooch_skip", r"\bbrooch\b"),  # blocked via HARD too
    ("earrings_for_round_face", r"round face|oval face|face shape"),
    ("jewellery_for_saree", r"jewellery for saree|jewelry for saree|necklace for saree"),
    ("jewellery_for_lehenga", r"lehenga"),
    ("jewellery_for_suit", r"for suit|salwar"),
    ("jewellery_care", r"how to clean|jewellery care|jewelry care|maintain gold"),
    ("gold_vs_diamond", r"gold vs|diamond vs|difference between"),
    ("certified_diamond", r"certified diamond|igi|gia"),
    ("lab_grown", r"lab grown|lab-grown|cvd|hpht"),
    ("solitaire_necklace", r"solitaire necklace|solitaire pendant"),
    ("solitaire_earrings", r"solitaire earring"),
    ("tennis_necklace", r"tennis necklace"),
    ("riviere", r"riviere"),
    ("lariat", r"lariat"),
    ("collar_necklace", r"collar necklace"),
    ("bib_necklace", r"bib necklace"),
    ("opera_necklace", r"opera length|opera necklace"),
    ("princess_cut", r"princess cut"),
    ("emerald_cut", r"emerald cut"),
    ("pear_cut", r"pear (cut|shape|diamond)"),
    ("halo_ring", r"halo ring"),
    ("eternity_ring", r"eternity"),
    ("signet_ring", r"signet"),
    ("midi_ring", r"midi ring"),
    ("thumb_ring", r"thumb ring"),
    ("adjustable_ring", r"adjustable ring|open ring"),
    ("adjustable_bracelet", r"adjustable bracelet"),
    ("openable_bangle", r"openable bangle|screw bangle"),
    ("kada_women_gold", r"gold kada for women|womens kada"),
    ("oxidised_earrings", r"oxidis.*earring"),
    ("oxidised_necklace", r"oxidis.*(necklace|set)"),
    ("temple_necklace", r"temple necklace"),
    ("temple_bangles", r"temple bangle"),
    ("antique_earrings", r"antique earring"),
    ("antique_necklace", r"antique necklace"),
    ("kundan_set", r"kundan set"),
    ("polki_set", r"polki set"),
    ("meenakari_earrings", r"meenakari earring|meena earring"),
    ("pearl_necklace", r"pearl necklace|pearl set"),
    ("pearl_bracelet", r"pearl bracelet"),
    ("ruby_jewellery", r"\bruby\b"),
    ("emerald_jewellery", r"\bemerald\b"),
    ("sapphire_jewellery", r"\bsapphire\b"),
    ("moissanite", r"moissanite"),
    ("cubic_zirconia", r"cubic zirconia|\bcz\b"),
    ("american_diamond", r"american diamond"),
    ("imitation_jewellery", r"imitation|artificial jewellery|fashion jewellery"),
    ("real_gold_tips", r"real gold|original gold"),
]

BLOCK_INTENTS = {
    "wedding_couple",
    "rakhi_sister",
    "karwa_wife",
    "gold_purity",
    "bangle_size",
    "wedding_girls_her",
    "sister_birthday",
}

HARD = re.compile(
    r"\b(caratlane|tanishq|kalyan|malabar|joyalukkas|bluestone|mmtc|pamp|digi ?gold|digital gold|"
    r"brooch|ms dhoni|sakharpuda|moradabad|photo|photos|image|images|pic|pics|video|wallpaper|"
    r"under\s*\d+|below\s*\d+|less than\s*\d+|rs\.?|rupees|price|rate today|emi|near me|store|showroom|branch|"
    r"buy online|amazon|flipkart|gift card|giftcard|check balance|retirement|"
    r"baby boy|baby girl|1 year baby|newborn baby|baby gift|godh bharai|"
    r"couple gift|gifts? for couples?|anniversary gifts? for (couple|friends)|"
    r"gift items for marriage|marriage below|lovers? day|"
    r"great valentines? gifts?|personalized gifts? for husband|"
    r"imitation|artificial jewellery|fashion jewellery|american diamond|cubic zirconia|\bcz\b|moissanite|"
    r"\bv ring\b|ring earrings|1 gram|3gm|3 gm|24k gold earrings mens)\b",
    re.I,
)
WISH = re.compile(r"\b(wish|wishes|quotes?|shayari|caption|status|\bmsg\b|msgs|greeting|greetings)\b", re.I)
GIFT = re.compile(r"\bgifts?\b", re.I)
JEWEL = re.compile(
    r"\b(earring|earrings|ring|rings|bangle|bangles|bracelet|bracelets|necklace|pendant|chain|"
    r"mangalsutra|anklet|payal|choker|jhumka|jhumki|tikka|coin|jewellery|jewelry|kada|nath|"
    r"polki|kundan|jadau|meenakari|solitaire|diamond|gold|platinum|silver|toe ring|bichiya|"
    r"hathphool|kamarband|bajuband|stud|hoop|charm|tennis)\b",
    re.I,
)


def intent_of(pk: str) -> str | None:
    n = norm(pk)
    for name, pat in RULES:
        if re.search(pat, n):
            return name
    if GIFT.search(pk) or JEWEL.search(pk):
        toks = tokens(pk)[:5]
        if len(toks) >= 2:
            return "sig:" + "_".join(toks)
    return None


def is_jewellery_gift_or_design(kw: str) -> bool:
    """New Week7 fills must be jewellery-shaped (not generic gift basket KWs)."""
    if JEWEL.search(kw):
        return True
    # Occasion gift only if clearly jewellery-gifting recipient we can own
    if GIFT.search(kw) and re.search(
        r"\b(jewellery|jewelry|gold|diamond|silver|pendant|necklace|earring|ring|bracelet|"
        r"bangle|chain|mangalsutra|coin)\b",
        kw,
        re.I,
    ):
        return True
    return False


def is_blog_worthy(kw: str, cluster: str = "", blog_fmt: str = "", page_fmt: str = "", theme: str = "") -> bool:
    if HARD.search(kw):
        return False
    if WISH.search(kw) and not GIFT.search(kw):
        return False
    # Expansion quality gate: skip generic gifts without jewellery signal
    if GIFT.search(kw) and not is_jewellery_gift_or_design(kw):
        return False
    edu_ok = bool(
        theme == "How-to / Education"
        or re.search(r"how to|size chart|gst|hallmark|purity|making charges|\b(9|14|18|22)\s*(k|kt|karat|carat)\b", kw, re.I)
        or "Educational" in (cluster or "")
        or blog_fmt in ("Guide / How-to", "Price / Buying guide")
    )
    if not is_jewellery_gift_or_design(kw) and not edu_ok:
        return False
    if blog_fmt in (
        "Festive / Wishes",
        "Listicle / Designs",
        "Guide / How-to",
        "Seasonal / Gifting",
        "Editorial / Trends",
        "Other blog",
        "Price / Buying guide",
    ):
        return True
    if page_fmt == "Blog":
        return True
    if cluster.startswith("4.") or cluster.startswith("8.") or "Educational" in cluster or "Festive" in cluster:
        return True
    if GIFT.search(kw):
        return is_jewellery_gift_or_design(kw)
    if JEWEL.search(kw) and re.search(
        r"design|types|for women|for men|for kids|gift|wear|set|latest|simple|traditional|modern|lightweight",
        kw,
        re.I,
    ):
        return True
    if theme in ("Gift Guides", "Design Listicles", "How-to / Education"):
        return True
    return False


def classify(kw: str) -> tuple[str, str]:
    if GIFT.search(kw) or re.search(
        r"diwali|christmas|valentine|friendship|akshaya|secret santa|engagement|house.?warming|rakhi|dhanteras|karwa|karva",
        kw,
        re.I,
    ):
        if re.search(
            r"diwali|christmas|valentine|friendship|akshaya|secret santa|engagement|house.?warming|rakhi|dhanteras|karwa|karva",
            kw,
            re.I,
        ):
            return "Gift Guides", "Seasonal / Gifting"
        return "Gift Guides", "Occasion/Gifting"
    if re.search(r"how to|size|gst|purity|hallmark|buy gold|making charges|karat|carat", kw, re.I):
        return "How-to / Education", "Buying Guide"
    return "Design Listicles", "Jewellery Design"


def year_for(kw: str) -> int:
    return 2027 if re.search(r"valentine|mother|father|baby|new mom|house.?warming", kw, re.I) else 2026


def slugify(kw: str, year: int) -> str:
    s = re.sub(r"[^a-z0-9\s]", "", kw.lower())
    s = re.sub(r"\b20\d{2}\b", "", s)
    s = re.sub(r"\s+", "-", s.strip()).strip("-")
    return f"{s}-{year}"[:90]


def main() -> None:
    os.chdir(ROOT)
    wb = load_workbook(SRC)
    ws = wb["Week 7"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    existing = []
    for r in range(2, ws.max_row + 1):
        row = {k: ws.cell(r, c).value for k, c in h.items() if k}
        try:
            rank = int(row.get("Priority Rank") or 0)
        except Exception:
            rank = 0
        # Freeze ranks 1..FREEZE_THROUGH; ranks above are rebuilt to reach TARGET.
        if rank and rank <= FREEZE_THROUGH:
            existing.append(row)
    existing = sorted(existing, key=lambda x: int(x.get("Priority Rank") or 0))[:FREEZE_THROUGH]
    print("existing Week7 frozen", len(existing), "target", TARGET)

    used_intents: set[str] = set()
    used_pks: list[str] = []
    used_exact: set[str] = set()
    for r in existing:
        pk = r.get("Primary Keyword") or ""
        ii = intent_of(pk)
        if ii:
            used_intents.add(ii)
        used_pks.append(pk)
        used_exact.add(norm(pk))
    used_intents |= BLOCK_INTENTS
    used_intents.add("wedding_couple")

    for sheet in ["Week 1-2", "Week 3-4", "Week 5", "Week 6"]:
        w = wb[sheet]
        hh = {w.cell(1, c).value: c for c in range(1, w.max_column + 1)}
        for r in range(2, w.max_row + 1):
            pk = w.cell(r, hh["Primary Keyword"]).value
            if not pk:
                continue
            used_exact.add(norm(pk))
            ii = intent_of(pk)
            if ii:
                used_intents.add(ii)
            used_pks.append(pk)
            sk = w.cell(r, hh.get("Supporting Keywords")).value or ""
            for p in re.split(r"\s*\|\s*", sk):
                if p.strip():
                    used_exact.add(norm(p))
                    ii = intent_of(p)
                    if ii:
                        used_intents.add(ii)

    print("blocked intents", len(used_intents))

    cands: list[dict] = []

    with open("output/Competitor_Master_Data.csv", newline="", encoding="utf-8", errors="replace") as f:
        for row in csv.DictReader(f):
            kw = (row.get("Keyword") or "").strip()
            if not kw:
                continue
            brand = row.get("Brand") or ""
            if brand.lower() == "bluestone":
                continue
            try:
                vol = int(float(row.get("Volume") or 0))
            except Exception:
                vol = 0
            if vol < 1000:
                continue
            kd_raw = row.get("KD")
            try:
                kd = float(kd_raw) if kd_raw not in (None, "") else None
            except Exception:
                kd = None
            cluster = row.get("Cluster") or ""
            blog_fmt = row.get("Blog Format") or ""
            page_fmt = row.get("Page Format") or ""
            url = row.get("URL") or ""
            if not is_blog_worthy(kw, cluster, blog_fmt, page_fmt):
                continue
            if page_fmt == "Category PLP" and not (
                GIFT.search(kw) or re.search(r"design|types|gift|for women|for men|latest|simple", kw, re.I)
            ):
                continue
            if page_fmt in ("Gold rate page", "Store page", "Homepage") and not GIFT.search(kw):
                continue
            if page_fmt == "Product PDP" and not re.search(r"design|gift|types|for women|for men", kw, re.I):
                continue
            if norm(kw) in used_exact:
                continue
            ii = intent_of(kw)
            if not ii or ii in used_intents:
                continue
            score = vol / (1 + (kd or 25) / 25)
            cands.append(
                {
                    "kw": kw,
                    "vol": vol,
                    "kd": kd,
                    "score": score,
                    "brand": brand,
                    "intent": ii,
                    "url": url,
                    "pos": row.get("Position"),
                    "source": "Competitor_Master_Data",
                }
            )

    print("competitor cands", len(cands), Counter(c["brand"] for c in cands).most_common())

    ws2 = wb["Keywords Scoring CL"]
    h2 = {ws2.cell(1, c).value: c for c in range(1, ws2.max_column + 1)}
    for r in range(2, ws2.max_row + 1):
        kw = (ws2.cell(r, h2["Keyword"]).value or "").strip()
        if not kw or HARD.search(kw):
            continue
        theme = ws2.cell(r, h2["Theme"]).value or ""
        blog = ws2.cell(r, h2["Blog Format"]).value or ""
        cluster = ws2.cell(r, h2["Cluster"]).value or ""
        try:
            vol = int(ws2.cell(r, h2["Volume"]).value or 0)
        except Exception:
            vol = 0
        if vol < 1000:
            continue
        if not is_blog_worthy(kw, cluster, blog, "", theme):
            continue
        if norm(kw) in used_exact:
            continue
        ii = intent_of(kw)
        if not ii or ii in used_intents:
            continue
        kd_raw = ws2.cell(r, h2["KD"]).value
        try:
            kd = float(kd_raw) if kd_raw is not None else None
        except Exception:
            kd = None
        score = float(ws2.cell(r, h2["Priority Score"]).value or 0) or (vol / (1 + (kd or 25) / 25))
        cands.append(
            {
                "kw": kw,
                "vol": vol,
                "kd": kd,
                "score": score,
                "brand": "ScoringCL",
                "intent": ii,
                "url": ws2.cell(r, h2["CaratLane URL"]).value or "",
                "pos": ws2.cell(r, h2["CL Position"]).value,
                "source": "Keywords Scoring CL",
            }
        )

    try:
        with open(
            "output/Jewellery_Keywords_Unique_All_Sources_by_Volume.csv",
            newline="",
            encoding="utf-8",
            errors="replace",
        ) as f:
            for row in csv.DictReader(f):
                kw = (row.get("Keyword") or "").strip()
                if not kw or HARD.search(kw):
                    continue
                try:
                    vol = int(float(row.get("Search Volume") or 0))
                except Exception:
                    vol = 0
                if vol < 1300:
                    continue
                if not is_blog_worthy(kw, theme="Design Listicles"):
                    continue
                if norm(kw) in used_exact:
                    continue
                ii = intent_of(kw)
                if not ii or ii in used_intents:
                    continue
                kd_raw = row.get("Keyword Difficulty")
                try:
                    kd = float(kd_raw) if kd_raw not in (None, "") else None
                except Exception:
                    kd = None
                cands.append(
                    {
                        "kw": kw,
                        "vol": vol,
                        "kd": kd,
                        "score": vol / (1 + (kd or 25) / 25),
                        "brand": "MergedJewellery",
                        "intent": ii,
                        "url": "",
                        "pos": None,
                        "source": "Jewellery_Keywords_Unique",
                    }
                )
    except Exception as e:
        print("jewellery file skip", e)

    print("total cands", len(cands))

    by_intent: dict[str, list[dict]] = defaultdict(list)
    for c in cands:
        by_intent[c["intent"]].append(c)

    picked = []
    for intent, items in by_intent.items():
        if intent in used_intents:
            continue
        items.sort(key=lambda x: (-x["vol"], -x["score"]))
        brands = set(i["brand"] for i in items)
        # Prefer jewellery-shaped keywords as primary (never generic gift hampers)
        jewel_items = [i for i in items if is_jewellery_gift_or_design(i["kw"])]
        if not jewel_items:
            continue
        primary = jewel_items[0]
        gifts = [i for i in jewel_items if GIFT.search(i["kw"])]
        if gifts and gifts[0]["vol"] >= primary["vol"] * 0.7:
            primary = gifts[0]
        primary = dict(primary)
        primary["brand_count"] = len(brands)
        primary["brands"] = ",".join(sorted(brands)[:6])
        primary["supports"] = [i["kw"] for i in items[1:8] if i["kw"] != primary["kw"]]
        picked.append(primary)

    picked.sort(key=lambda p: (-(1 if GIFT.search(p["kw"]) else 0), -p["brand_count"], -p["vol"], -p["score"]))

    need = TARGET - len(existing)
    added = []

    def append_pick(p: dict, note_tag: str) -> None:
        year = year_for(p["kw"])
        theme, cat = classify(p["kw"])
        seen = {norm(p["kw"])}
        merg = []
        for s in p.get("supports") or []:
            if norm(s) in seen:
                continue
            seen.add(norm(s))
            merg.append(s)
        row = {
            "Priority Rank": None,
            "Week": "Week 7",
            "Month Plan Bucket": f"Committed - Gift Guides Post Rank 10 (intent-unique, multi-brand, target {TARGET})",
            "Source": f"Multi-brand data ({p['source']}; brands={p['brands']})",
            "Type": "Validated Variant Page",
            "Category Fit": cat,
            "Theme": theme,
            "Action": "New",
            "Primary Keyword": p["kw"],
            "Article Title/Angle": re.sub(r"\b20\d{2}\b", str(year), p["kw"], flags=re.I),
            "Suggested URL Slug": slugify(p["kw"], year),
            "Bluestone Blog URL": None,
            "Supporting Keywords": " | ".join(merg[:8]) if merg else p["kw"],
            "Keyword Count": 1 + min(len(merg), 8),
            "Volume": p["vol"],
            "KD": p["kd"],
            "Priority Score": round(float(p["score"]), 4) if p["score"] else None,
            "In CaratLane Export": "Yes" if "CaratLane" in p["brands"] else "No",
            "CL Position": p["pos"] if p.get("brand") == "CaratLane" else None,
            "Bluestone Position": None,
            "CaratLane URL": p["url"] if "caratlane" in (p.get("url") or "").lower() else None,
            "Semrush Page": p["kw"],
            "Execution Note": (
                f"Week7 multi-brand fill ({note_tag}). intent={p['intent']}; brands={p['brands']}; "
                f"from {p['source']}. One URL only. Live WP duplicate-intent gate required. No prices."
            ),
            "_intent": p["intent"],
            "_brands": p["brands"],
        }
        added.append(row)
        used_intents.add(p["intent"])
        used_exact.add(norm(p["kw"]))
        used_pks.append(p["kw"])

    for p in picked:
        if len(added) >= need:
            break
        if p["intent"] in used_intents or norm(p["kw"]) in used_exact:
            continue
        if any(stem_overlap(p["kw"], prev) >= 0.45 for prev in used_pks):
            continue
        if not is_jewellery_gift_or_design(p["kw"]) and not re.search(
            r"how to|size|gst|hallmark|purity|making charges|karat|carat", p["kw"], re.I
        ):
            continue
        if p["intent"].startswith("sig:") and (p["vol"] < 1400 or len(tokens(p["kw"])) < 2):
            continue
        append_pick(p, "strict")

    print("after strict added", len(added), "total", len(existing) + len(added))

    if len(existing) + len(added) < TARGET:
        extra_pool = []
        with open("output/Competitor_Master_Data.csv", newline="", encoding="utf-8", errors="replace") as f:
            for row in csv.DictReader(f):
                kw = (row.get("Keyword") or "").strip()
                brand = row.get("Brand") or ""
                if not kw or brand.lower() == "bluestone" or HARD.search(kw):
                    continue
                if WISH.search(kw) and not GIFT.search(kw):
                    continue
                if not JEWEL.search(kw):
                    continue
                try:
                    vol = int(float(row.get("Volume") or 0))
                except Exception:
                    vol = 0
                if vol < 1200:
                    continue
                if not re.search(
                    r"design|types|for women|for men|for kids|for girls|gift|latest|simple|traditional|wedding|daily wear|party wear|lightweight|set\b",
                    kw,
                    re.I,
                ):
                    continue
                ii = intent_of(kw)
                if not ii:
                    toks = tokens(kw)[:4]
                    if len(toks) < 2:
                        continue
                    ii = "sig:" + "_".join(toks)
                if ii in used_intents or norm(kw) in used_exact:
                    continue
                extra_pool.append(
                    {
                        "kw": kw,
                        "vol": vol,
                        "kd": row.get("KD"),
                        "brand": brand,
                        "intent": ii,
                        "url": row.get("URL") or "",
                        "source": "Competitor_Master_loose",
                        "score": vol,
                        "pos": row.get("Position"),
                    }
                )
        by2: dict[str, list[dict]] = defaultdict(list)
        for e in extra_pool:
            by2[e["intent"]].append(e)
        loose = []
        for intent, items in by2.items():
            items.sort(key=lambda x: -x["vol"])
            brands = sorted(set(i["brand"] for i in items))
            primary = items[0]
            loose.append(
                {
                    **primary,
                    "brands": ",".join(brands),
                    "brand_count": len(brands),
                    "supports": [i["kw"] for i in items[1:6]],
                }
            )
        loose.sort(key=lambda x: (-x["brand_count"], -x["vol"]))
        for p in loose:
            if len(existing) + len(added) >= TARGET:
                break
            if p["intent"] in used_intents:
                continue
            if any(stem_overlap(p["kw"], prev) >= 0.45 for prev in used_pks):
                continue
            try:
                p["kd"] = float(p["kd"]) if p["kd"] not in (None, "") else None
            except Exception:
                p["kd"] = None
            p["score"] = p["vol"] / (1 + (p["kd"] or 25) / 25)
            append_pick(p, "loose-design")
        print("after loose added", len(added), "total", len(existing) + len(added))

    # Pass 3: lower volume floor for more multi-brand design/gift leftovers
    if len(existing) + len(added) < TARGET:
        extra2 = []
        with open("output/Competitor_Master_Data.csv", newline="", encoding="utf-8", errors="replace") as f:
            for row in csv.DictReader(f):
                kw = (row.get("Keyword") or "").strip()
                brand = row.get("Brand") or ""
                if not kw or brand.lower() == "bluestone" or HARD.search(kw):
                    continue
                if not is_jewellery_gift_or_design(kw):
                    continue
                if WISH.search(kw) and not GIFT.search(kw):
                    continue
                try:
                    vol = int(float(row.get("Volume") or 0))
                except Exception:
                    vol = 0
                if vol < 700:
                    continue
                if not re.search(
                    r"design|types|for women|for men|for kids|for girls|gift|latest|simple|"
                    r"traditional|wedding|daily wear|party wear|lightweight|set\b|bangle|earring|"
                    r"necklace|pendant|bracelet|ring|chain|mangalsutra|anklet|jhumka|kada",
                    kw,
                    re.I,
                ):
                    continue
                ii = intent_of(kw)
                if not ii:
                    toks = tokens(kw)[:4]
                    if len(toks) < 2:
                        continue
                    ii = "sig:" + "_".join(toks)
                if ii in used_intents or norm(kw) in used_exact:
                    continue
                extra2.append(
                    {
                        "kw": kw,
                        "vol": vol,
                        "kd": row.get("KD"),
                        "brand": brand,
                        "intent": ii,
                        "url": row.get("URL") or "",
                        "source": "Competitor_Master_v700",
                        "score": vol,
                        "pos": row.get("Position"),
                    }
                )
        by3: dict[str, list[dict]] = defaultdict(list)
        for e in extra2:
            by3[e["intent"]].append(e)
        loose2 = []
        for intent, items in by3.items():
            items.sort(key=lambda x: -x["vol"])
            jewel_items = [i for i in items if is_jewellery_gift_or_design(i["kw"])]
            if not jewel_items:
                continue
            brands = sorted(set(i["brand"] for i in items))
            primary = jewel_items[0]
            loose2.append(
                {
                    **primary,
                    "brands": ",".join(brands),
                    "brand_count": len(brands),
                    "supports": [i["kw"] for i in items[1:6]],
                }
            )
        loose2.sort(key=lambda x: (-x["brand_count"], -x["vol"]))
        for p in loose2:
            if len(existing) + len(added) >= TARGET:
                break
            if p["intent"] in used_intents:
                continue
            if any(stem_overlap(p["kw"], prev) >= 0.45 for prev in used_pks):
                continue
            try:
                p["kd"] = float(p["kd"]) if p["kd"] not in (None, "") else None
            except Exception:
                p["kd"] = None
            p["score"] = p["vol"] / (1 + (p["kd"] or 25) / 25)
            append_pick(p, "v700-fill")
        print("after v700 added", len(added), "total", len(existing) + len(added))

    final = (existing + added)[:TARGET]
    for i, r in enumerate(final, 1):
        r["Priority Rank"] = i
        r["Week"] = "Week 7"

    # Collapse any intent collisions (including within previously frozen ranks)
    # Keep highest-volume row per intent; refill vacated slots from leftover picks.
    ib_rows: dict[str, list[dict]] = defaultdict(list)
    for r in final:
        ii = intent_of(r.get("Primary Keyword") or "") or ("raw:" + norm(r.get("Primary Keyword") or ""))
        ib_rows[ii].append(r)
    keep: list[dict] = []
    dropped = 0
    for ii, group in ib_rows.items():
        group.sort(key=lambda x: (-(int(x.get("Volume") or 0)), int(x.get("Priority Rank") or 999)))
        keep.append(group[0])
        dropped += len(group) - 1
    # Also drop HARD junk that may sit in frozen ranks
    cleaned = []
    for r in keep:
        pk = r.get("Primary Keyword") or ""
        if HARD.search(pk) and r.get("Action") == "New" and int(r.get("Priority Rank") or 0) > 9:
            dropped += 1
            continue
        cleaned.append(r)
    keep = cleaned
    print("after intent collapse kept", len(keep), "dropped", dropped)

    used_intents = set()
    used_exact = set()
    used_pks = []
    for r in keep:
        pk = r.get("Primary Keyword") or ""
        ii = intent_of(pk)
        if ii:
            used_intents.add(ii)
        used_exact.add(norm(pk))
        used_pks.append(pk)
    used_intents |= BLOCK_INTENTS

    # Rebuild refill pool from leftover picked intents not used
    refill_added = []
    for p in picked:
        if len(keep) + len(refill_added) >= TARGET:
            break
        if p["intent"] in used_intents or norm(p["kw"]) in used_exact:
            continue
        if any(stem_overlap(p["kw"], prev) >= 0.45 for prev in used_pks):
            continue
        if not is_jewellery_gift_or_design(p["kw"]):
            continue
        if HARD.search(p["kw"]):
            continue
        year = year_for(p["kw"])
        theme, cat = classify(p["kw"])
        seen = {norm(p["kw"])}
        merg = []
        for s in p.get("supports") or []:
            if norm(s) in seen:
                continue
            seen.add(norm(s))
            merg.append(s)
        row = {
            "Priority Rank": None,
            "Week": "Week 7",
            "Month Plan Bucket": f"Committed - Gift Guides Post Rank 10 (intent-unique, multi-brand, target {TARGET})",
            "Source": f"Multi-brand data ({p['source']}; brands={p.get('brands','')})",
            "Type": "Validated Variant Page",
            "Category Fit": cat,
            "Theme": theme,
            "Action": "New",
            "Primary Keyword": p["kw"],
            "Article Title/Angle": re.sub(r"\b20\d{2}\b", str(year), p["kw"], flags=re.I),
            "Suggested URL Slug": slugify(p["kw"], year),
            "Bluestone Blog URL": None,
            "Supporting Keywords": " | ".join(merg[:8]) if merg else p["kw"],
            "Keyword Count": 1 + min(len(merg), 8),
            "Volume": p["vol"],
            "KD": p["kd"],
            "Priority Score": round(float(p["score"]), 4) if p.get("score") else None,
            "In CaratLane Export": "Yes" if "CaratLane" in str(p.get("brands", "")) else "No",
            "CL Position": p["pos"] if p.get("brand") == "CaratLane" else None,
            "Bluestone Position": None,
            "CaratLane URL": p["url"] if "caratlane" in (p.get("url") or "").lower() else None,
            "Semrush Page": p["kw"],
            "Execution Note": (
                f"Week7 multi-brand refill after intent collapse. intent={p['intent']}; "
                f"brands={p.get('brands')}; from {p['source']}. One URL only."
            ),
            "_intent": p["intent"],
            "_brands": p.get("brands"),
        }
        refill_added.append(row)
        used_intents.add(p["intent"])
        used_exact.add(norm(p["kw"]))
        used_pks.append(p["kw"])

    # Prefer Done/Skip continuity rows first, then by prior rank, then volume
    def sort_key(r):
        act = r.get("Action") or "New"
        pri = {"Done": 0, "Skip-Existing": 1, "New": 2}.get(act, 3)
        return (pri, int(r.get("Priority Rank") or 999), -(int(r.get("Volume") or 0)))

    final = sorted(keep + refill_added, key=sort_key)[:TARGET]
    # If still short, pull from competitor again at vol>=500
    if len(final) < TARGET:
        print("still short", len(final), "mining vol>=500")
        pool500 = []
        with open("output/Competitor_Master_Data.csv", newline="", encoding="utf-8", errors="replace") as f:
            for row in csv.DictReader(f):
                kw = (row.get("Keyword") or "").strip()
                brand = row.get("Brand") or ""
                if not kw or brand.lower() == "bluestone" or HARD.search(kw):
                    continue
                if not is_jewellery_gift_or_design(kw):
                    continue
                try:
                    vol = int(float(row.get("Volume") or 0))
                except Exception:
                    vol = 0
                if vol < 500:
                    continue
                ii = intent_of(kw)
                if not ii:
                    toks = tokens(kw)[:4]
                    if len(toks) < 2:
                        continue
                    ii = "sig:" + "_".join(toks)
                if ii in used_intents or norm(kw) in used_exact:
                    continue
                if any(stem_overlap(kw, prev) >= 0.45 for prev in used_pks):
                    continue
                pool500.append((vol, kw, brand, ii, row.get("KD"), row.get("URL") or "", row.get("Position")))
        pool500.sort(reverse=True)
        by_i: dict[str, list] = defaultdict(list)
        for item in pool500:
            by_i[item[3]].append(item)
        for ii, items in sorted(by_i.items(), key=lambda kv: -kv[1][0][0]):
            if len(final) >= TARGET:
                break
            vol, kw, brand, intent, kd, url, pos = items[0]
            brands = ",".join(sorted({x[2] for x in items}))
            try:
                kd_f = float(kd) if kd not in (None, "") else None
            except Exception:
                kd_f = None
            year = year_for(kw)
            theme, cat = classify(kw)
            final.append(
                {
                    "Priority Rank": None,
                    "Week": "Week 7",
                    "Month Plan Bucket": f"Committed - Gift Guides Post Rank 10 (intent-unique, multi-brand, target {TARGET})",
                    "Source": f"Multi-brand data (Competitor_Master_v500; brands={brands})",
                    "Type": "Validated Variant Page",
                    "Category Fit": cat,
                    "Theme": theme,
                    "Action": "New",
                    "Primary Keyword": kw,
                    "Article Title/Angle": re.sub(r"\b20\d{2}\b", str(year), kw, flags=re.I),
                    "Suggested URL Slug": slugify(kw, year),
                    "Bluestone Blog URL": None,
                    "Supporting Keywords": " | ".join(x[1] for x in items[1:6]) or kw,
                    "Keyword Count": min(len(items), 6),
                    "Volume": vol,
                    "KD": kd_f,
                    "Priority Score": round(vol / (1 + (kd_f or 25) / 25), 4),
                    "In CaratLane Export": "Yes" if "CaratLane" in brands else "No",
                    "CL Position": pos if brand == "CaratLane" else None,
                    "Bluestone Position": None,
                    "CaratLane URL": url if "caratlane" in url.lower() else None,
                    "Semrush Page": kw,
                    "Execution Note": f"Week7 v500 fill. intent={intent}; brands={brands}. One URL only.",
                    "_intent": intent,
                    "_brands": brands,
                }
            )
            used_intents.add(intent)
            used_exact.add(norm(kw))
            used_pks.append(kw)

    final = final[:TARGET]
    for i, r in enumerate(final, 1):
        r["Priority Rank"] = i
        r["Week"] = "Week 7"

    ib: dict[str, list[str]] = defaultdict(list)
    for r in final:
        ii = intent_of(r.get("Primary Keyword") or "") or ("raw:" + norm(r.get("Primary Keyword") or ""))
        ib[ii].append(r.get("Primary Keyword") or "")
    dups = {k: v for k, v in ib.items() if len(v) > 1}
    print("FINAL", len(final), "dups", len(dups))
    if dups:
        print("dup sample", list(dups.items())[:8])

    # Audit: rows beyond prior 150 sheet length
    prior_n = min(FREEZE_THROUGH, len(existing))
    added_rows = final[prior_n:]
    existing_for_audit = prior_n
    added_for_audit = len(added_rows)

    headers = [
        "Priority Rank",
        "Week",
        "Month Plan Bucket",
        "Source",
        "Type",
        "Category Fit",
        "Theme",
        "Action",
        "Primary Keyword",
        "Article Title/Angle",
        "Suggested URL Slug",
        "Bluestone Blog URL",
        "Supporting Keywords",
        "Keyword Count",
        "Volume",
        "KD",
        "Priority Score",
        "In CaratLane Export",
        "CL Position",
        "Bluestone Position",
        "CaratLane URL",
        "Semrush Page",
        "Execution Note",
    ]
    colors = {"New": "C6EFCE", "Done": "BDD7EE", "Skip-Existing": "F4B183"}

    def write(path: str) -> None:
        wb2 = load_workbook(path)
        if "Week 7" in wb2.sheetnames:
            del wb2["Week 7"]
        idx = wb2.sheetnames.index("Week 6") + 1 if "Week 6" in wb2.sheetnames else None
        ws2 = wb2.create_sheet("Week 7", idx)
        hf = PatternFill("solid", "1B4F72")
        hfont = Font(color="FFFFFF", bold=True)
        thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        for c, name in enumerate(headers, 1):
            cell = ws2.cell(1, c, name)
            cell.fill = hf
            cell.font = hfont
        for r in final:
            for c, name in enumerate(headers, 1):
                cell = ws2.cell(r["Priority Rank"] + 1, c, r.get(name))
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if name == "Action":
                    cell.fill = PatternFill("solid", colors.get(r.get("Action"), "C6EFCE"))
                if r["Priority Rank"] >= 10 and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "FFF2CC")
                if r["Priority Rank"] > FREEZE_THROUGH and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "D9EAD3")
                if FREEZE_THROUGH >= 74 and 74 <= r["Priority Rank"] <= FREEZE_THROUGH and name == "Priority Rank":
                    cell.fill = PatternFill("solid", "CFE2F3")
        ws2.freeze_panes = "A2"
        ws2.auto_filter.ref = f"A1:W{len(final) + 1}"
        widths = {
            "A": 8,
            "B": 10,
            "C": 52,
            "D": 48,
            "E": 20,
            "F": 18,
            "G": 18,
            "H": 14,
            "I": 42,
            "J": 42,
            "K": 42,
            "L": 40,
            "M": 45,
            "N": 10,
            "O": 10,
            "P": 8,
            "Q": 12,
            "R": 12,
            "S": 10,
            "T": 12,
            "U": 28,
            "V": 28,
            "W": 55,
        }
        for col, w in widths.items():
            ws2.column_dimensions[col].width = w
        wb2.save(path)

    for path in PATHS:
        write(path)
        print("saved", path)

    out_dir = "final seo generation context/output"
    os.makedirs(out_dir, exist_ok=True)
    with open(f"{out_dir}/Week7_Gift_Guides_PostRank10.csv", "w", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "rank",
                "action",
                "intent",
                "primary",
                "volume",
                "kd",
                "score",
                "theme",
                "brands_or_source",
                "slug",
                "url",
            ],
        )
        w.writeheader()
        for r in final:
            src = r.get("Source") or ""
            brands = ""
            m = re.search(r"brands=([^)]+)", src)
            if m:
                brands = m.group(1)
            w.writerow(
                {
                    "rank": r["Priority Rank"],
                    "action": r.get("Action"),
                    "intent": intent_of(r.get("Primary Keyword") or "") or "",
                    "primary": r.get("Primary Keyword"),
                    "volume": r.get("Volume"),
                    "kd": r.get("KD"),
                    "score": r.get("Priority Score"),
                    "theme": r.get("Theme"),
                    "brands_or_source": brands or src[:60],
                    "slug": r.get("Suggested URL Slug"),
                    "url": r.get("Bluestone Blog URL"),
                }
            )

    audit = {
        "total": len(final),
        "existing_frozen": existing_for_audit,
        "added": added_for_audit,
        "duplicate_intent_count": len(dups),
        "added_theme_mix": dict(Counter(r.get("Theme") for r in added_rows)),
        "actions": dict(Counter(r.get("Action") for r in final)),
        "themes": dict(Counter(r.get("Theme") for r in final)),
        "added_sample": [
            {
                "rank": r["Priority Rank"],
                "primary": r["Primary Keyword"],
                "intent": r.get("_intent") or intent_of(r.get("Primary Keyword") or ""),
                "brands": r.get("_brands"),
                "vol": r["Volume"],
            }
            for r in added_rows[:50]
        ],
    }
    json.dump(audit, open(f"{out_dir}/week7_expand_{TARGET}_audit.json", "w"), indent=2)
    print("DONE total", len(final))
    print("actions", Counter(r.get("Action") for r in final))
    print("themes", Counter(r.get("Theme") for r in final))
    start = max(FREEZE_THROUGH, len(final) - 30)
    print(f"Ranks {FREEZE_THROUGH + 1}-{min(FREEZE_THROUGH + 25, len(final))}:")
    for r in final[FREEZE_THROUGH : FREEZE_THROUGH + 25]:
        print(f"{r['Priority Rank']:3}. vol={r.get('Volume')} {r['Primary Keyword']}")


if __name__ == "__main__":
    main()
