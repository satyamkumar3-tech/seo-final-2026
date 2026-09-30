#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 82: Rose Gold Bangle Buying Guide."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 82
PRIMARY = "rose gold bangle"
SLUG = "rose-gold-bangle-2026"
BLOG_URL = "https://blog.bluestone.com/rose-gold-bangle-2026/"
POST_ID = "40381"
CAROUSEL_MEDIA = "40375,40376,40377,40378,40379,40380"

def update_status_csv(type3_str: str, visible_words: int):
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": type3_str,
                        "lines": "",
                        "visible_words": str(visible_words),
                        "notes": "2026-09-27T15:10:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": type3_str,
            "lines": "",
            "visible_words": str(visible_words),
            "notes": "2026-09-27T15:10:00Z | SERP intelligence pipeline v1"
        })

    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-27"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BIDG0393O37 The Estrella Oval Bangle",
        "flatlay": "BIPS0003O06 The Channing Bangle",
        "lifestyle": "BISM0003O14 The Muricelle Bangle"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "study-desk",
        "note": "rose gold bangle The Channing Bangle on study desk with brass calipers and loupe"
    })

    ranks = data.setdefault("ranks", {})
    ranks[str(RANK)] = {
        "rank": str(RANK),
        "primary": PRIMARY,
        "slug": SLUG,
        "hero_sku": "BIDG0393O37",
        "flatlay_sku": "BIPS0003O06",
        "lifestyle_sku": "BISM0003O14",
        "flatlay_setting": "study-desk"
    }

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    print(f"Updated {rotation_file} with Rank {RANK} product rotation")

def update_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print(f"Workbook {xlsx_path} not found, skipping.")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    if "Week 9" not in wb.sheetnames:
        print("Week 9 sheet not found in workbook, skipping.")
        return

    ws = wb["Week 9"]
    headers = [cell.value for cell in ws[1]]

    rank_col = None
    url_col = None
    action_col = None
    note_col = None

    for idx, h in enumerate(headers, 1):
        if h in ["Priority Rank", "Rank", "rank"]:
            rank_col = idx
        elif h == "Bluestone Blog URL":
            url_col = idx
        elif h in ["Action", "action"]:
            action_col = idx
        elif h in ["Execution Note", "execution_note"]:
            note_col = idx

    for row_idx in range(2, ws.max_row + 1):
        r_val = ws.cell(row=row_idx, column=rank_col).value
        if str(r_val) == str(RANK):
            if url_col:
                ws.cell(row=row_idx, column=url_col, value=BLOG_URL)
            if action_col:
                ws.cell(row=row_idx, column=action_col, value="Done")
            if note_col:
                note_text = f"Published 2026-09-27. WP ID {POST_ID}. Engine: buying_guide. Author Satyam (270271337). 3D Coverflow + Type 3 trio verified."
                ws.cell(row=row_idx, column=note_col, value=note_text)
            print(f"Updated Excel row {row_idx} for Rank {RANK}: URL, Action=Done, Note")
            break

    wb.save(xlsx_path)
    print("Workbook saved successfully!")

def main():
    type3_uploaded_path = ROOT / "output/Week9_Rank82_type3_uploaded_media.json"
    type3_str = ""
    if type3_uploaded_path.exists():
        with open(type3_uploaded_path) as f:
            t3 = json.load(f)
            type3_str = f"{t3['hero']['media_id']},{t3['flatlay']['media_id']},{t3['lifestyle']['media_id']}"

    qa_path = ROOT / "output/week9_rank82_live_qa.json"
    words = 3563
    if qa_path.exists():
        with open(qa_path) as f:
            qa = json.load(f)
            words = qa.get("visible_words", words)

    update_status_csv(type3_str, words)
    update_product_rotation()
    update_workbook()
    print("All status updates completed successfully!")

if __name__ == "__main__":
    main()
