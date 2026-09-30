#!/usr/bin/env python3
"""Update status files and Excel workbook for Week 8 Rank 178: Finger Rings for Girls 2026."""
import os
import csv
import json
from pathlib import Path
import openpyxl

ROOT = Path("/Users/satyamkumar/Downloads/seo final 2026")
RANK = 178
SLUG = "finger-rings-for-girls-2026"
PRIMARY = "finger rings for girls"
POST_ID = 39077
URL = "https://blog.bluestone.com/finger-rings-for-girls-2026/"
WORD_COUNT = "3510"

# 1. Update output/Week8_Blog_Queue_status.csv
status_csv_path = ROOT / "output/Week8_Blog_Queue_status.csv"
fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            if str(r.get("rank", "")).strip() != str(RANK):
                rows.append(r)

type3_str = "hero 39078; flatlay 39079; lifestyle 39080"
note_text = f"Published 2026-09-16 via buying_guide engine. Comprehensive guide to finger rings for girls, 18K/14K gold purity, BIS 3-mark hallmarking, ring finger placement & symbolism, band ergonomics & safe settings, 6-card 3D coverflow carousel (39071-39076), 3 Type 3 images (The Anya Ring, The Quinn Ring, The Rafia Ring), flatlay setting marble-vanity, author Satyam (270271337), passed live QA."

rank178_row = {
    "rank": str(RANK),
    "primary": PRIMARY,
    "slug": SLUG,
    "blog_url": URL,
    "wp_post_id": str(POST_ID),
    "status": "published",
    "carousel_media": "39071-39076",
    "type3_media": type3_str,
    "lines": "425",
    "visible_words": WORD_COUNT,
    "notes": note_text
}
rows.append(rank178_row)
rows.sort(key=lambda x: int(x["rank"]) if x["rank"].isdigit() else 9999)

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/Week8_Blog_Queue.csv if present
queue_csv_path = ROOT / "output/Week8_Blog_Queue.csv"
if queue_csv_path.exists():
    q_rows = []
    with open(queue_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        q_fields = reader.fieldnames
        for r in reader:
            if str(r.get("rank", "")).strip() == str(RANK):
                r["action"] = "Done"
            q_rows.append(r)
    with open(queue_csv_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=q_fields)
        writer.writeheader()
        writer.writerows(q_rows)
    print(f"Updated {queue_csv_path}")

# 3. Update output/product_rotation.json
rot_path = ROOT / "output/product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {"recent_flatlay_settings": [], "recent_skus": []}

recent_trios = rot_data.setdefault("recent_type3_trios", [])
# Remove existing entry for this rank if any
recent_trios = [t for t in recent_trios if isinstance(t, dict) and t.get("rank") != f"Week8 Rank {RANK}" and t.get("rank") != RANK]
recent_trios.append({
    "rank": f"Week8 Rank {RANK}",
    "hero": "BIAR0097R04 The Anya Ring",
    "flatlay": "BIAR0097R16 The Quinn Ring",
    "lifestyle": "BIAB0503R03 The Rafia Ring"
})
rot_data["recent_type3_trios"] = recent_trios

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings = [s for s in recent_settings if (not isinstance(s, dict)) or (s.get("rank") != f"Week8 Rank {RANK}" and s.get("rank") != RANK)]
recent_settings.append({
    "rank": f"Week8 Rank {RANK}",
    "setting": "marble-vanity",
    "note": "finger rings for girls The Quinn Ring"
})
rot_data["recent_flatlay_settings"] = recent_settings
rot_data["recent_flatlay_settings"] = recent_settings

recent_skus = rot_data.setdefault("recent_skus", [])
recent_skus = [s for s in recent_skus if s.get("rank") != f"Week8 Rank {RANK}"]
for sku_info in [
    {"rank": f"Week8 Rank {RANK}", "slot": "hero", "code": "BIAR0097R04", "name": "The Anya Ring"},
    {"rank": f"Week8 Rank {RANK}", "slot": "flatlay", "code": "BIAR0097R16", "name": "The Quinn Ring"},
    {"rank": f"Week8 Rank {RANK}", "slot": "lifestyle", "code": "BIAB0503R03", "name": "The Rafia Ring"}
]:
    recent_skus.append(sku_info)
rot_data["recent_skus"] = recent_skus

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2)
print(f"Updated {rot_path}")

# 4. Update SEO Strategy 2026.xlsx
xlsx_target = ROOT / "SEO Strategy 2026.xlsx"
real_xlsx = xlsx_target.resolve()
wb = openpyxl.load_workbook(real_xlsx)
if "Week 8" in wb.sheetnames:
    ws = wb["Week 8"]
    headers = [cell.value for cell in ws[1]]
    rank_col = headers.index("Priority Rank") + 1
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1

    target_row_idx = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=rank_col).value
        if str(val).strip() == str(RANK):
            target_row_idx = r
            break

    if target_row_idx:
        ws.cell(row=target_row_idx, column=url_col, value=URL)
        ws.cell(row=target_row_idx, column=note_col, value=f"Published 2026-09-16. WP Post ID: {POST_ID}. Author: Satyam (270271337). Engine: buying_guide. Categories: Gold (554493348), Jewellery Problem & Solution (554493465). Type 3 trio: The Anya Ring (hero), The Quinn Ring (flatlay), The Rafia Ring (lifestyle). Flatlay setting: marble-vanity. Carousel: 39071-39076. Live QA passed with {WORD_COUNT} words, verified BIS HUID/18K/14K standards, single H1, 3D Coverflow carousel, and og:image.")
        wb.save(real_xlsx)
        print(f"Updated {real_xlsx} sheet Week 8 row {target_row_idx}")
    else:
        print(f"WARNING: Rank {RANK} row not found in Week 8 sheet")
else:
    print("WARNING: 'Week 8' sheet not found in workbook")
