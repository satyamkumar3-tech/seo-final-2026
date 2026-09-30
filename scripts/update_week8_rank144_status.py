#!/usr/bin/env python3
"""Update status files and Excel workbook for Week 8 Rank 144: Purple Earrings."""
import os
import csv
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

# 1. Update output/Week8_Blog_Queue_status.csv
status_csv_path = ROOT / "output/Week8_Blog_Queue_status.csv"
fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            if str(r.get("rank", "")).strip() != "144":
                rows.append(r)

qa_path = ROOT / "output/Week8_Rank144_PurpleEarrings_qa_results.json"
word_count = "2976"
if qa_path.exists():
    with open(qa_path) as f:
        qa_data = json.load(f)
        word_count = str(qa_data.get("word_count", "2976"))

type3_path = ROOT / "output/Week8_Rank144_PurpleEarrings_type3_media.json"
type3_str = "hero 38717; flatlay 38718; lifestyle 38719"
if type3_path.exists():
    with open(type3_path) as f:
        t3_data = json.load(f)
        h_id = t3_data.get("hero", {}).get("id", "38717")
        f_id = t3_data.get("flatlay", {}).get("id", "38718")
        l_id = t3_data.get("lifestyle", {}).get("id", "38719")
        type3_str = f"hero {h_id}; flatlay {f_id}; lifestyle {l_id}"

rank144_row = {
    "rank": "144",
    "primary": "purple earrings",
    "slug": "purple-earrings-2026",
    "blog_url": "https://blog.bluestone.com/purple-earrings-2026/",
    "wp_post_id": "38716",
    "status": "published",
    "carousel_media": "38710-38715",
    "type3_media": type3_str,
    "lines": "410",
    "visible_words": word_count,
    "notes": f"Published 2026-09-11 via buying_guide engine. Comprehensive guide to purple earrings, natural amethyst/tanzanite/sapphire, 18Kt/14Kt gold pairings, BIS HUID hallmarking, ethnic and western styling frameworks, 6-card 3D coverflow carousel (38710-38715), 3 Type 3 images (The Aleena Huggie Earrings, The Nettile Huggie Earrings, The Vicky Hoop Earrings), flatlay setting windowsill-daylight, author Satyam (270271337), passed live QA."
}
rows.append(rank144_row)
rows.sort(key=lambda x: int(x["rank"]))

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/product_rotation.json
rot_path = ROOT / "output/product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {"recent_flatlay_settings": [], "recent_skus": []}

recent_trios = rot_data.setdefault("recent_type3_trios", [])
recent_trios.append({
    "rank": "Week8 Rank 144",
    "hero": "BIIP0279S08 The Aleena Huggie Earrings",
    "flatlay": "BIPN0880H218 The Nettile Huggie Earrings",
    "lifestyle": "BIIP0427H16 The Vicky Hoop Earrings"
})

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings.append({
    "rank": "Week8 Rank 144",
    "setting": "windowsill-daylight",
    "note": "purple earrings The Nettile Huggie Earrings"
})

recent_skus = rot_data.setdefault("recent_skus", [])
for sku_info in [
    {"rank": "Week8 Rank 144", "slot": "hero", "code": "BIIP0279S08", "name": "The Aleena Huggie Earrings"},
    {"rank": "Week8 Rank 144", "slot": "flatlay", "code": "BIPN0880H218", "name": "The Nettile Huggie Earrings"},
    {"rank": "Week8 Rank 144", "slot": "lifestyle", "code": "BIIP0427H16", "name": "The Vicky Hoop Earrings"}
]:
    recent_skus.append(sku_info)

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2)
print(f"Updated {rot_path}")

# 3. Update SEO Strategy 2026.xlsx
xlsx_target = ROOT / "SEO Strategy 2026.xlsx"
real_xlsx = xlsx_target.resolve()
wb = openpyxl.load_workbook(real_xlsx)
if "Week 8" in wb.sheetnames:
    ws = wb["Week 8"]
    headers = [cell.value for cell in ws[1]]
    rank_col = headers.index("Priority Rank") + 1
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1
    
    target_row_idx = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=rank_col).value
        if str(val).strip() == "144":
            target_row_idx = r
            break
            
    if target_row_idx:
        ws.cell(row=target_row_idx, column=url_col, value="https://blog.bluestone.com/purple-earrings-2026/")
        ws.cell(row=target_row_idx, column=note_col, value=f"Published 2026-09-11. WP Post ID: 38716. Author: Satyam (270271337). Engine: buying_guide. Categories: Gold (554493348), Jewellery Problem & Solution (554493465). Type 3 trio: The Aleena Huggie Earrings (hero), The Nettile Huggie Earrings (flatlay), The Vicky Hoop Earrings (lifestyle). Flatlay setting: windowsill-daylight. Carousel: 38710-38715. Live QA passed with {word_count} words, verified BIS HUID/18Kt/14Kt standards, single H1, and og:image.")
        wb.save(real_xlsx)
        print(f"Updated {real_xlsx} sheet Week 8 row {target_row_idx}")
    else:
        print("WARNING: Rank 144 row not found in Week 8 sheet")
else:
    print("WARNING: 'Week 8' sheet not found in workbook")
