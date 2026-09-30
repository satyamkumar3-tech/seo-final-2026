#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 73."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 73
PRIMARY = "marriage ring finger"
SLUG = "marriage-ring-finger-2026"
BLOG_URL = "https://blog.bluestone.com/marriage-ring-finger-2026/"
POST_ID = "40275"
CAROUSEL_MEDIA = "40269,40270,40271,40272,40273,40274"
TYPE3_MEDIA = "40276,40277,40278"
VISIBLE_WORDS = "2954"
NOTES = "Published 2026-09-26. WP#40275. Author Satyam 270271337. Type 3 trio: BISL0851R28 (The Jasper Band For Him) / BIAR0097R04 (The Anya Ring) / BIPM0017R18 (The Malibu Ring). Flatlay: study-desk. 6 carousel cards live. BIS hallmarking HUID & ring sizing ergonomics verified. Zero dash/price errors."

def update_status_csv():
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": TYPE3_MEDIA,
                        "lines": "",
                        "visible_words": VISIBLE_WORDS,
                        "notes": "2026-09-26T16:32:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": TYPE3_MEDIA,
            "lines": "",
            "visible_words": VISIBLE_WORDS,
            "notes": "2026-09-26T16:32:00Z | SERP intelligence pipeline v1"
        })

    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-26"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BISL0851R28 The Jasper Band For Him",
        "flatlay": "BIAR0097R04 The Anya Ring",
        "lifestyle": "BIPM0017R18 The Malibu Ring"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "study-desk",
        "note": "marriage ring finger The Anya Ring on dark walnut study desk with brass ring sizer gauge, antique loupe, closed leather notebook, and vintage fountain pen"
    })

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"Updated {rotation_file} with Rank {RANK} rotation and flatlay setting")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print("Excel workbook not found!")
        return

    try:
        wb = openpyxl.load_workbook(xlsx_path)
        if "Week 9" not in wb.sheetnames:
            print("Sheet Week 9 not found!")
            return

        ws = wb["Week 9"]
        headers = [c.value for c in ws[1]]
        url_col = headers.index("Bluestone Blog URL") + 1
        note_col = headers.index("Execution Note") + 1
        rank_col = headers.index("Priority Rank") + 1

        updated_row = None
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=rank_col).value) == str(RANK):
                ws.cell(row=row, column=url_col, value=BLOG_URL)
                ws.cell(row=row, column=note_col, value=NOTES)
                updated_row = row
                break

        if updated_row:
            wb.save(xlsx_path)
            print(f"Updated Excel workbook row {updated_row} for Rank {RANK}")
        else:
            print(f"Row for Rank {RANK} not found in Excel workbook")
    except Exception as e:
        print(f"Warning: could not update Excel workbook: {e}")

if __name__ == "__main__":
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()
