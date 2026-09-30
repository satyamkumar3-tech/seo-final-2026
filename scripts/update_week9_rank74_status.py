#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 74: Neelam Stone Ring Buying Guide."""
import csv
import json
import os
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

def update_status_csv():
    csv_path = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    rows = []
    header = None
    rank_found = False

    new_row = {
        "rank": "74",
        "primary": "neelam stone ring",
        "slug": "neelam-stone-ring-2026",
        "blog_url": "https://blog.bluestone.com/neelam-stone-ring-2026/",
        "wp_post_id": "40286",
        "status": "published",
        "carousel_media": "40280,40281,40282,40283,40284,40285",
        "type3_media": "40287,40288,40289",
        "lines": "",
        "visible_words": "3443",
        "notes": "Completed by buying_guide with checkpointed WordPress, media, Type 3 generation, social metadata verification, and passed live QA."
    }

    if csv_path.exists():
        with open(csv_path, mode="r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            header = reader.fieldnames
            for row in reader:
                if str(row.get("rank")) == "74":
                    rows.append(new_row)
                    rank_found = True
                else:
                    rows.append(row)

    if not rank_found:
        rows.append(new_row)

    if not header:
        header = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

    with open(csv_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=header)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {csv_path} with Rank 74.")

def update_product_rotation():
    rot_path = ROOT / "output" / "product_rotation.json"
    if not rot_path.exists():
        print("product_rotation.json not found!")
        return

    with open(rot_path, mode="r", encoding="utf-8") as f:
        data = json.load(f)

    # Upsert post
    post_entry = {
        "rank": "Week9 Rank 74",
        "slug": "neelam-stone-ring-2026",
        "flatlay_setting": "cafe-tray",
        "skus": {
            "hero": "BIAR0097R07",
            "flatlay": "BINS0639R18",
            "lifestyle": "BIAB0503R03"
        },
        "names": {
            "hero": "The Liza Ring",
            "flatlay": "The Gigi Ring",
            "lifestyle": "The Rafia Ring"
        },
        "carousel_skus": [
            "BIAR0097R07",
            "BINS0639R18",
            "BISL0851R28",
            "BIJP0993R123",
            "BISV0910R24",
            "BINS0639R11"
        ]
    }

    posts = data.get("posts", [])
    posts = [p for p in posts if p.get("rank") != "Week9 Rank 74" and p.get("slug") != "neelam-stone-ring-2026"]
    posts.append(post_entry)
    data["posts"] = posts

    # Update recent_flatlay_settings
    recent_settings = data.get("recent_flatlay_settings", [])
    recent_settings.append({
        "rank": "Week9 Rank 74",
        "setting": "cafe-tray",
        "note": "neelam stone ring The Gigi Ring on stoneware cafe tray with ceramic saucer and brass loupe"
    })
    data["recent_flatlay_settings"] = recent_settings

    with open(rot_path, mode="w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    print(f"Updated {rot_path} with Rank 74.")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print("Workbook not found!")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    if "Week 9" not in wb.sheetnames:
        print("Sheet 'Week 9' not in workbook!")
        return

    ws = wb["Week 9"]
    headers = [cell.value for cell in ws[1]]
    url_col_idx = headers.index("Bluestone Blog URL") + 1 if "Bluestone Blog URL" in headers else None
    note_col_idx = headers.index("Execution Note") + 1 if "Execution Note" in headers else None
    rank_col_idx = headers.index("Priority Rank") + 1 if "Priority Rank" in headers else 1

    updated = False
    for row in range(2, ws.max_row + 1):
        rank_val = ws.cell(row=row, column=rank_col_idx).value
        if str(rank_val) == "74":
            if url_col_idx:
                ws.cell(row=row, column=url_col_idx).value = "https://blog.bluestone.com/neelam-stone-ring-2026/"
            if note_col_idx:
                ws.cell(row=row, column=note_col_idx).value = "Published post 40286 on 2026-09-26. 3443 words, 11 H2s, 3D coverflow carousel, Type 3 hero/flatlay/lifestyle (nano_banana_pro)."
            updated = True
            print(f"Row {row} (Rank 74) updated in Sheet 'Week 9'.")
            break

    if updated:
        wb.save(xlsx_path)
        print("Saved SEO Strategy 2026.xlsx successfully.")
    else:
        print("Rank 74 not found in Week 9 sheet.")

def main():
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()

if __name__ == "__main__":
    main()
