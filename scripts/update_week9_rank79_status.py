#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 79: Modern Gold Long Necklace Designs."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 79
PRIMARY = "modern gold long necklace designs"
SLUG = "modern-gold-long-necklace-designs-2026"
BLOG_URL = "https://blog.bluestone.com/modern-gold-long-necklace-designs-2026/"
POST_ID = "40347"
CAROUSEL_MEDIA = "40341,40342,40343,40344,40345,40346"
TYPE3_MEDIA = "40348,40349,40350"
VISIBLE_WORDS = "3723"
NOTES = "Published 2026-09-26. WP#40347. Author Satyam 270271337. Type 3 trio: BINS0780C08 (The Ninetta Mangalsutra Necklace) / BIPN0987N07 (The Rapett Evil Eye Charm Necklace) / BIMA0780C53 (The Casma Mangalsutra). Flatlay: gift-wrapping-station. 6 carousel cards live. BIS hallmarking HUID & net gold weight billing verified. Zero dash/price errors."

def update_status_csv():
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": TYPE3_MEDIA,
                        "lines": "",
                        "visible_words": VISIBLE_WORDS,
                        "notes": "2026-09-26T23:05:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": TYPE3_MEDIA,
            "lines": "",
            "visible_words": VISIBLE_WORDS,
            "notes": "2026-09-26T23:05:00Z | SERP intelligence pipeline v1"
        })

    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-26"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BINS0780C08 The Ninetta Mangalsutra Necklace",
        "flatlay": "BIPN0987N07 The Rapett Evil Eye Charm Necklace",
        "lifestyle": "BIMA0780C53 The Casma Mangalsutra"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "gift-wrapping-station",
        "note": f"{PRIMARY} The Rapett Evil Eye Charm Necklace on oatmeal linen runner with brass calipers and folded sage ribbon"
    })

    ranks = data.setdefault("ranks", {})
    ranks[str(RANK)] = {
        "primary": PRIMARY,
        "slug": SLUG,
        "type3_skus": ["BINS0780C08", "BIPN0987N07", "BIMA0780C53"],
        "flatlay_setting": "gift-wrapping-station",
        "carousel_skus": [
            "BIAV0987N78",
            "BISL0819N09",
            "BIPN0987N07",
            "BVPJ0935C06",
            "BINS0780C08",
            "BIMA1081C01"
        ]
    }

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    print(f"Updated {rotation_file} with Rank {RANK} rotation details")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print(f"Warning: {xlsx_path} not found")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    sheet = wb["Week 9"]

    headers = [cell.value for cell in sheet[1]]
    url_col = headers.index("Bluestone Blog URL") + 1 if "Bluestone Blog URL" in headers else None
    slug_col = headers.index("Suggested URL Slug") + 1 if "Suggested URL Slug" in headers else None
    note_col = headers.index("Execution Note") + 1 if "Execution Note" in headers else None
    action_col = headers.index("Action") + 1 if "Action" in headers else None

    row_found = False
    for r in range(2, sheet.max_row + 1):
        cell_val = sheet.cell(row=r, column=1).value
        if str(cell_val) == str(RANK):
            if url_col:
                sheet.cell(row=r, column=url_col, value=BLOG_URL)
            if slug_col:
                sheet.cell(row=r, column=slug_col, value=SLUG)
            if note_col:
                sheet.cell(row=r, column=note_col, value=NOTES)
            if action_col:
                sheet.cell(row=r, column=action_col, value="Done")
            row_found = True
            print(f"Updated XLSX row {r} for Rank {RANK}")
            break

    if row_found:
        wb.save(xlsx_path)
        print(f"Saved changes to {xlsx_path}")
    else:
        print(f"Could not find Rank {RANK} row in {xlsx_path}")

if __name__ == "__main__":
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()
