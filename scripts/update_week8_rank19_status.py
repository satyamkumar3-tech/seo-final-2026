#!/usr/bin/env python3
"""Update status files and Excel workbook for Week 8 Rank 19: Engagement Rings for Couples."""
import os
import csv
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

# 1. Update output/Week8_Blog_Queue_status.csv
status_csv_path = ROOT / "output/Week8_Blog_Queue_status.csv"
fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]

rows = []
if status_csv_path.exists():
    with open(status_csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            if r["rank"] != "19":
                rows.append(r)

rank19_row = {
    "rank": "19",
    "primary": "engagement rings for couples",
    "slug": "engagement-rings-for-couples-2026",
    "blog_url": "https://blog.bluestone.com/engagement-rings-for-couples-2026/",
    "wp_post_id": "37205",
    "status": "published",
    "carousel_media": "37199-37204",
    "type3_media": "hero 37206; flatlay 37207; lifestyle 37208",
    "lines": "407",
    "visible_words": "3333",
    "notes": "Published 2026-09-01 via buying_guide engine. Comprehensive guide to matching couple engagement rings, 18K/14K gold durability, 4Cs diamond selection, BIS HUID hallmarking, SGL/IGI/GIA certification, sizing precision, 6-card carousel, 3 Type 3 images (The Liza Ring, The Interlink Band Ring, The Jasper Band For Him), passed live QA."
}
rows.append(rank19_row)
rows.sort(key=lambda x: int(x["rank"]))

with open(status_csv_path, mode="w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print(f"Updated {status_csv_path}")

# 2. Update output/product_rotation.json
rot_path = ROOT / "output/product_rotation.json"
if rot_path.exists():
    with open(rot_path, "r", encoding="utf-8") as f:
        rot_data = json.load(f)
else:
    rot_data = {"recent_flatlay_settings": [], "recent_skus": []}

recent_trios = rot_data.setdefault("recent_type3_trios", [])
recent_trios.append({
    "rank": "Week8 Rank 19",
    "hero": "BIAR0097R07 The Liza Ring",
    "flatlay": "BISV0910R24 The Interlink Band Ring",
    "lifestyle": "BISL0851R28 The Jasper Band For Him"
})

recent_settings = rot_data.setdefault("recent_flatlay_settings", [])
recent_settings.append({
    "rank": "Week8 Rank 19",
    "setting": "marble-vanity",
    "note": "couple engagement rings The Interlink Band Ring"
})

recent_skus = rot_data.setdefault("recent_skus", [])
for sku_info in [
    {"rank": "Week8 Rank 19", "slot": "hero", "code": "BIAR0097R07", "name": "The Liza Ring"},
    {"rank": "Week8 Rank 19", "slot": "flatlay", "code": "BISV0910R24", "name": "The Interlink Band Ring"},
    {"rank": "Week8 Rank 19", "slot": "lifestyle", "code": "BISL0851R28", "name": "The Jasper Band For Him"}
]:
    recent_skus.append(sku_info)

with open(rot_path, "w", encoding="utf-8") as f:
    json.dump(rot_data, f, indent=2)
print(f"Updated {rot_path}")

# 3. Update SEO Strategy 2026.xlsx
xlsx_target = ROOT / "SEO Strategy 2026.xlsx"
real_xlsx = xlsx_target.resolve()
wb = openpyxl.load_workbook(real_xlsx)
if "Week 8" in wb.sheetnames:
    ws = wb["Week 8"]
    headers = [cell.value for cell in ws[1]]
    rank_col = headers.index("Priority Rank") + 1
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1
    
    target_row_idx = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=rank_col).value
        if str(val).strip() == "19":
            target_row_idx = r
            break
            
    if target_row_idx:
        ws.cell(row=target_row_idx, column=url_col, value="https://blog.bluestone.com/engagement-rings-for-couples-2026/")
        ws.cell(row=target_row_idx, column=note_col, value="Published 2026-09-01. WP Post ID: 37205. Author: Satyam (270271337). Engine: buying_guide. Categories: Wedding Jewellery (554493443), Rings (554493418), Gold (554493348), Jewellery Problem & Solution (554493465), Gift (554493424). Type 3 trio: The Liza Ring (hero 37206), The Interlink Band Ring (flatlay 37207), The Jasper Band For Him (lifestyle 37208). Flatlay setting: marble-vanity. Carousel: 37199-37204. Live QA passed with 3333 words, verified 4Cs/HUID/SGL/IGI/GIA standards, single H1, and og:image.")
        wb.save(real_xlsx)
        print(f"Updated {real_xlsx} sheet Week 8 row {target_row_idx}")
    else:
        print("WARNING: Rank 19 row not found in Week 8 sheet")
else:
    print("WARNING: 'Week 8' sheet not found in workbook")
