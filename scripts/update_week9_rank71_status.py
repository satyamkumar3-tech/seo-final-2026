#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update status files for Week 9 Rank 71."""
import csv
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

RANK = 71
PRIMARY = "light weight gold earrings design"
SLUG = "light-weight-gold-earrings-design-2026"
BLOG_URL = "https://blog.bluestone.com/light-weight-gold-earrings-design-2026/"
POST_ID = "40253"
CAROUSEL_MEDIA = "40247,40248,40249,40250,40251,40252"
TYPE3_MEDIA = "40254,40255,40256"
VISIBLE_WORDS = "3213"
NOTES = "Published 2026-09-25. WP#40253. Author Satyam 270271337. Type 3 trio: BIIP0427H16 / BIJP0686H03 / BINK0363H03. Flatlay: windowsill-daylight. 6 carousel cards live. BIS hallmarking HUID & net weight billing verified. Zero dash/price errors."

def update_status_csv():
    status_file = ROOT / "output" / "Week9_Blog_Queue_status.csv"
    fieldnames = ["rank", "primary", "slug", "blog_url", "wp_post_id", "status", "carousel_media", "type3_media", "lines", "visible_words", "notes"]
    rows = []
    found = False

    if status_file.exists():
        with open(status_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if str(r.get("rank")) == str(RANK):
                    rows.append({
                        "rank": str(RANK),
                        "primary": PRIMARY,
                        "slug": SLUG,
                        "blog_url": BLOG_URL,
                        "wp_post_id": POST_ID,
                        "status": "Done",
                        "carousel_media": CAROUSEL_MEDIA,
                        "type3_media": TYPE3_MEDIA,
                        "lines": "",
                        "visible_words": VISIBLE_WORDS,
                        "notes": "2026-09-25T11:56:00Z | SERP intelligence pipeline v1"
                    })
                    found = True
                else:
                    rows.append(r)

    if not found:
        rows.append({
            "rank": str(RANK),
            "primary": PRIMARY,
            "slug": SLUG,
            "blog_url": BLOG_URL,
            "wp_post_id": POST_ID,
            "status": "Done",
            "carousel_media": CAROUSEL_MEDIA,
            "type3_media": TYPE3_MEDIA,
            "lines": "",
            "visible_words": VISIBLE_WORDS,
            "notes": "2026-09-25T11:56:00Z | SERP intelligence pipeline v1"
        })

    # Sort by rank numeric
    def rank_key(row):
        try:
            return int(row["rank"])
        except:
            return 9999

    rows.sort(key=rank_key)

    with open(status_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Updated {status_file} with Rank {RANK} status Done")

def update_product_rotation():
    rotation_file = ROOT / "output" / "product_rotation.json"
    if not rotation_file.exists():
        data = {"recent_type3_trios": [], "recent_flatlay_settings": []}
    else:
        with open(rotation_file, "r", encoding="utf-8") as f:
            data = json.load(f)

    data["updated"] = "2026-09-25"

    trios = data.setdefault("recent_type3_trios", [])
    trios.append({
        "rank": f"Week9 Rank {RANK}",
        "hero": "BIIP0427H16 The Vicky Hoop Earrings",
        "flatlay": "BIJP0686H03 The Faliha Purse Hoop Earrings",
        "lifestyle": "BINK0363H03 The Skein Hoop Earrings"
    })

    settings = data.setdefault("recent_flatlay_settings", [])
    settings.append({
        "rank": f"Week9 Rank {RANK}",
        "setting": "windowsill-daylight",
        "note": "light weight gold earrings design The Faliha Purse Hoop Earrings on lime-washed oak windowsill with ceramic dish and brass loupe"
    })

    with open(rotation_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"Updated {rotation_file} with Rank {RANK} rotation and flatlay setting")

def update_excel_workbook():
    xlsx_path = ROOT / "SEO Strategy 2026.xlsx"
    if not xlsx_path.exists():
        print("Excel workbook not found!")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    if "Week 9" not in wb.sheetnames:
        print("Sheet Week 9 not found!")
        return

    ws = wb["Week 9"]
    headers = [c.value for c in ws[1]]
    url_col = headers.index("Bluestone Blog URL") + 1
    note_col = headers.index("Execution Note") + 1
    rank_col = headers.index("Priority Rank") + 1

    updated_row = None
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=rank_col).value == RANK:
            ws.cell(row=row, column=url_col, value=BLOG_URL)
            ws.cell(row=row, column=note_col, value=NOTES)
            updated_row = row
            break

    if updated_row:
        wb.save(xlsx_path)
        print(f"Updated Excel workbook row {updated_row} for Rank {RANK}")
    else:
        print(f"Row for Rank {RANK} not found in Excel workbook")

if __name__ == "__main__":
    update_status_csv()
    update_product_rotation()
    update_excel_workbook()
